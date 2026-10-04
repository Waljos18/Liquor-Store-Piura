"""Genera Retrospectiva.docx a partir del contenido del formato de retrospectiva."""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


def set_cell_shading(cell, fill_hex: str) -> None:
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill_hex)
    cell._tc.get_or_add_tcPr().append(shading)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], header_fill: str = "D9E2F3") -> None:
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    hdr_cells = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].text = h
        set_cell_shading(hdr_cells[i], header_fill)
        for p in hdr_cells[i].paragraphs:
            for r in p.runs:
                r.bold = True
    for r_idx, row_data in enumerate(rows):
        row_cells = table.rows[r_idx + 1].cells
        for c_idx, text in enumerate(row_data):
            row_cells[c_idx].text = text
    doc.add_paragraph()


def main() -> None:
    out = Path(__file__).resolve().parent.parent / "Retrospectiva.docx"
    doc = Document()

    title = doc.add_heading("Formato de retrospectiva de Sprint", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    sub = doc.add_heading("Sistema Web y Aplicación Móvil – Licorería Chilalo Shot", level=2)
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER

    p = doc.add_paragraph()
    p.add_run("Metodología: ").bold = True
    p.add_run("Scrum — reunión al cierre de cada sprint (duración sugerida: 30–45 min).")
    p = doc.add_paragraph()
    p.add_run("Objetivo: ").bold = True
    p.add_run("Identificar qué funcionó, qué ajustar y qué acciones concretas llevar al siguiente sprint.")

    doc.add_heading("Cómo usar esta plantilla", level=1)
    for i, item in enumerate(
        [
            "Copia la sección Plantilla vacía al final de cada sprint y rellénala.",
            "Si el equipo es una sola persona (TAA), indícalo en «Tipo de retrospectiva».",
            "Las acciones deben ser específicas y revisables en el planning del siguiente sprint.",
            "Opcional: enlaza este documento con el Sprint Review correspondiente (Integrador 2/SPRINT_REVIEW_N.md).",
        ],
        start=1,
    ):
        doc.add_paragraph(f"{i}. {item}", style="List Number")

    doc.add_heading("Plantilla vacía (copiar desde aquí)", level=1)

    doc.add_heading("Retrospectiva del Sprint [número]", level=2)
    doc.add_heading("Información general", level=3)
    add_table(
        doc,
        ["Campo", "Contenido"],
        [
            ["Sprint", "Sprint _"],
            ["Período", "Semanas _ al _"],
            ["Objetivo del sprint", "_"],
            ["Fecha de la retrospectiva", "_"],
            ["Facilita", "_"],
            ["Asistentes", "_"],
            ["Tipo de retrospectiva", "Individual / Equipo"],
        ],
    )

    doc.add_heading("¿Qué salió bien?", level=3)
    doc.add_paragraph(
        "Listar logros, prácticas útiles, decisiones acertadas, herramientas que ayudaron."
    ).italic = True
    for _ in range(3):
        doc.add_paragraph(style="List Bullet")

    doc.add_heading("¿Qué se puede mejorar?", level=3)
    doc.add_paragraph(
        "Problemas, retrasos, malas estimaciones, deuda técnica, comunicación, riesgos no vistos."
    ).italic = True
    for _ in range(3):
        doc.add_paragraph(style="List Bullet")

    doc.add_heading("Acciones para el siguiente sprint", level=3)
    doc.add_paragraph("Cada acción debe ser concreta (qué, cómo, cuándo revisar).").italic = True
    add_table(
        doc,
        ["#", "Acción", "Responsable", "Revisión"],
        [["1", "", "", ""], ["2", "", "", ""], ["3", "", "", ""]],
    )

    doc.add_heading("Elementos que pasan al backlog del siguiente sprint (opcional)", level=3)
    add_table(
        doc,
        ["#", "Elemento", "Origen (plan / cliente / retrospectiva)"],
        [["1", "", ""], ["2", "", ""]],
    )

    doc.add_heading("Conclusión breve", level=3)
    doc.add_paragraph(
        "Párrafo de cierre: cumplimiento del objetivo del sprint y enfoque para el siguiente."
    ).italic = True
    doc.add_paragraph()

    foot = doc.add_paragraph()
    foot.add_run(
        "Trabajo Académico Aplicado (TAA) — Instituto de Educación Superior, Piura, Perú."
    ).italic = True

    doc.add_page_break()

    doc.add_heading("Ejemplo completado — Sprint 2 (referencia)", level=1)
    doc.add_paragraph(
        "Este bloque muestra el formato ya rellenado; puedes borrarlo o sustituirlo por tu sprint."
    ).italic = True

    doc.add_heading("Información general", level=3)
    add_table(
        doc,
        ["Campo", "Contenido"],
        [
            ["Sprint", "Sprint 2"],
            ["Período", "Semanas 5 al 9"],
            [
                "Objetivo del sprint",
                "POS web, facturación SUNAT (pruebas), inventario, compras, caja, gastos, devoluciones, crédito y clientes",
            ],
            ["Fecha de la retrospectiva", "(fin de semana 9)"],
            ["Facilita", "Alumno desarrollador"],
            ["Asistentes", "Alumno desarrollador"],
            ["Tipo de retrospectiva", "Individual"],
        ],
    )

    doc.add_heading("¿Qué salió bien?", level=3)
    bullets = [
        "Objetivo cumplido: operación digital con POS, stock automático y comprobantes en ambiente de pruebas.",
        "La base del Sprint 1 (JWT, Flyway, capas) evitó rediseños grandes.",
        "Inventario con alertas, compras, proveedores y mermas; caja, gastos, devoluciones y crédito cerraron el flujo operativo.",
        "Demo al cliente: venta → inventario → comprobante → caja → crédito con pago parcial.",
    ]
    for b in bullets:
        doc.add_paragraph(b, style="List Bullet")

    doc.add_heading("¿Qué se puede mejorar?", level=3)
    for b in [
        "Integración SUNAT/OSE: documentar incidencias (series, XML, correlativos).",
        "Pago mixto y casos borde: más tiempo del estimado; conviene trocear la historia antes del sprint.",
        "Poco margen para tests automatizados nuevos; priorizar tests en servicios críticos en el Sprint 3.",
    ]:
        doc.add_paragraph(b, style="List Bullet")

    doc.add_heading("Acciones para el siguiente sprint", level=3)
    add_table(
        doc,
        ["#", "Acción", "Responsable", "Revisión"],
        [
            [
                "1",
                "Partir historias complejas (promociones, IA, Android) en tareas pequeñas con criterios claros",
                "Desarrollador",
                "Planning Sprint 3",
            ],
            [
                "2",
                "Reservar colchón para integraciones externas e IA",
                "Desarrollador",
                "Inicio Sprint 3",
            ],
            [
                "3",
                "Aumentar tests en VentaService, facturación e inventario",
                "Desarrollador",
                "Durante Sprint 3",
            ],
        ],
    )

    doc.add_heading("Conclusión breve", level=3)
    doc.add_paragraph(
        "El Sprint 2 cerró según plan operativo; el Sprint 3 concentrará promociones, fidelización, reportes, IA, Android y despliegue."
    )

    doc.add_paragraph()
    p2 = doc.add_paragraph()
    p2.add_run(
        "Versión plantilla: 1.0 — compatible con Integrador 2/SPRINT_REVIEW_*.md y Retrospectiva_Sprint2.md."
    ).italic = True

    doc.save(out)
    print(f"Guardado: {out}")


if __name__ == "__main__":
    main()
