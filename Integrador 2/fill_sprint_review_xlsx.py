# -*- coding: utf-8 -*-
"""
Genera SPRINT_REVIEW_N_CHILALO_SHOT.xlsx desde FORMATO SPRINT REVIEW.xlsx (Hoja3).
Uso: python fill_sprint_review_xlsx.py 1|2|3
"""
import argparse
from pathlib import Path

import openpyxl

TEMPLATE = Path(r"c:\Users\Usuario\Downloads\FORMATO SPRINT REVIEW.xlsx")
SHEET_NAME = "Hoja3"

# (orden, id_hu, texto_hu, [(id_tarea, tarea, est, prioridad, status, procedencia, obs), ...])

SPRINT_1 = {
    "file": "SPRINT_REVIEW_1_CHILALO_SHOT.xlsx",
    "goal": (
        "Entregar la base técnica del sistema para la licorería Chilalo Shot: modelo de datos, "
        "autenticación JWT con roles (Administrador/Vendedor), recuperación de contraseña, gestión de "
        "usuarios, categorías y productos con búsqueda/filtros, configuración del negocio (RUC, razón social) "
        "y migraciones Flyway (V1–V4), dejando lista la arquitectura para el POS y la facturación SUNAT en el Sprint 2."
    ),
    "blocks": [
        (
            1,
            "HU-01",
            (
                "Como equipo de desarrollo, quiero dejar definidos el diseño de datos, la arquitectura y los requisitos "
                "para construir el sistema sobre una base consistente."
            ),
            [
                (
                    "T0001",
                    "Diseño del modelo de base de datos (MER): productos, categorías, usuarios, roles, proveedores, clientes.",
                    8,
                    "MUST",
                    "APROBADO",
                    None,
                    None,
                ),
                ("T0002", "Definición de arquitectura: backend REST Spring Boot + frontend React + PostgreSQL + app Android Kotlin.", 5, "MUST", "APROBADO", None, None),
                ("T0003", "Configuración del repositorio Git y carpetas backend / frontend / android.", 3, "MUST", "APROBADO", None, None),
                ("T0004", "Mockups de interfaces principales (POS, productos, inventario, dashboard) aprobados por el cliente (sem. 2).", 5, "MUST", "APROBADO", None, None),
                ("T0005", "Levantamiento de requerimientos con el propietario de Chilalo Shot (sem. 1).", 3, "MUST", "APROBADO", None, None),
            ],
        ),
        (
            2,
            "HU-02",
            (
                "Como usuario del sistema, quiero iniciar sesión de forma segura, recuperar mi contraseña y que el "
                "administrador gestione usuarios para controlar el acceso al negocio."
            ),
            [
                ("T0006", "Autenticación JWT: roles Administrador y Vendedor; contraseñas con BCrypt; tokens acceso/refresh.", 8, "MUST", "APROBADO", None, None),
                ("T0007", "Recuperación de contraseña por correo electrónico (enlace con token de seguridad).", 5, "MUST", "APROBADO", None, None),
                ("T0008", "Gestión de usuarios: crear, editar, activar/desactivar (solo Administrador).", 5, "MUST", "APROBADO", None, None),
            ],
        ),
        (
            3,
            "HU-03",
            (
                "Como administrador, quiero gestionar categorías y productos (con búsqueda y filtros) para mantener "
                "actualizado el catálogo de la licorería."
            ),
            [
                ("T0009", "Gestión de categorías: registro, edición y listado.", 3, "MUST", "APROBADO", None, None),
                (
                    "T0010",
                    "Gestión de productos: nombre, precios, categoría, código de barras, stocks, vencimiento, imagen; alta/edición/baja lógica.",
                    8,
                    "MUST",
                    "APROBADO",
                    None,
                    "10 productos reales cargados como prueba",
                ),
                ("T0011", "Búsqueda y filtrado de productos (nombre, código, categoría, activo).", 5, "MUST", "APROBADO", None, None),
            ],
        ),
        (
            4,
            "HU-04",
            (
                "Como administrador, quiero configurar los datos del negocio y contar con migraciones versionadas de BD "
                "para desplegar el esquema de forma ordenada."
            ),
            [
                ("T0012", "Configuración general del sistema: RUC, razón social, dirección, datos para futura facturación SUNAT.", 5, "MUST", "APROBADO", None, None),
                ("T0013", "Migraciones automáticas Flyway (V1 a V4) ejecutadas al iniciar el backend.", 3, "MUST", "APROBADO", None, None),
            ],
        ),
        (
            5,
            "HU-05",
            (
                "Como desarrollador, quiero adelantar bases de la app móvil y su conectividad para reducir riesgo en sprints posteriores "
                "(tareas adicionales acordadas en la revisión)."
            ),
            [
                ("T0014", "Configuración Android: network_security_config.xml para desarrollo (HTTP/certificados).", 2, "SHOULD", "APROBADO", None, "Extra Sprint 1"),
                ("T0015", "Estructura base app Android (Kotlin, Jetpack Compose, Hilt, Retrofit) y pantalla de login funcional.", 5, "SHOULD", "APROBADO", None, "Adelanto Sprint 3"),
            ],
        ),
    ],
    "summary": [
        ("Sprint", "Sprint 1"),
        ("Período", "Semanas 1–4"),
        ("Fecha revisión", "Fin de semana 4"),
        ("Cumplimiento backlog", "100% (13/13 tareas + 2 adicionales)"),
        ("Horas aproximadas", "~80 h"),
        ("Incidente principal", "CORS Windows (resuelto en sem. 1)"),
        ("Feedback cliente destacado", "Pantalla productos clara; pedido: ver stock en lista → backlog Sprint 2"),
    ],
}

SPRINT_2 = {
    "file": "SPRINT_REVIEW_2_CHILALO_SHOT.xlsx",
    "goal": (
        "Meta Sprint 2 (condensada): operación diaria digital — POS web con formas de pago, facturación electrónica SUNAT "
        "(sandbox), inventario y compras, caja, gastos, devoluciones, crédito/fiado y clientes (semanas 5–9)."
    ),
    "blocks": [
        (
            1,
            "HU-06",
            (
                "Como vendedor, quiero usar un POS web ágil para registrar ventas con carrito, descuentos, varias formas de pago "
                "y consultar el historial, viendo disponibilidad de stock en tiempo real."
            ),
            [
                ("T0201", "POS: búsqueda rápida de productos por nombre y código de barras; listado con stock visible.", 8, "MUST", "APROBADO", None, None),
                ("T0202", "Carrito de venta: cantidades editables, descuentos (porcentaje o monto) y totales automáticos.", 8, "MUST", "APROBADO", None, None),
                ("T0203", "Formas de pago: efectivo, tarjeta, Yape, Plin, transferencia, mixto y crédito/fiado.", 8, "MUST", "APROBADO", None, None),
                ("T0204", "Cálculo de vuelto y validaciones básicas antes de confirmar la venta.", 5, "MUST", "APROBADO", None, None),
                ("T0205", "Historial de ventas con filtros por fecha, forma de pago y estado.", 5, "MUST", "APROBADO", None, None),
                ("T0206", "Integración POS con inventario: bloqueo/agotado visible al vender (feedback Sprint 1).", 5, "MUST", "APROBADO", None, "Stock en lista"),
            ],
        ),
        (
            2,
            "HU-07",
            (
                "Como administrador, quiero emitir boletas y facturas electrónicas cumpliendo la normativa SUNAT desde el sistema."
            ),
            [
                ("T0207", "Integración backend con API SUNAT / OSE en ambiente de pruebas (sandbox).", 8, "MUST", "APROBADO", None, None),
                ("T0208", "Emisión de boletas electrónicas (serie B001) con XML y PDF.", 8, "MUST", "APROBADO", None, None),
                ("T0209", "Emisión de facturas electrónicas (serie F001) cuando corresponda.", 8, "MUST", "APROBADO", None, None),
                ("T0210", "Almacenamiento y consulta de estado de comprobantes; manejo de errores de envío.", 5, "MUST", "APROBADO", None, None),
                ("T0211", "Resumen diario de boletas (RCB) según flujo requerido.", 5, "MUST", "APROBADO", None, None),
            ],
        ),
        (
            3,
            "HU-08",
            (
                "Como administrador, quiero controlar inventario y abastecimiento: alertas, movimientos, compras, proveedores y mermas."
            ),
            [
                ("T0212", "Descuento automático de stock al confirmar ventas; consistencia con detalle de venta.", 5, "MUST", "APROBADO", None, None),
                ("T0213", "Alertas de stock bajo respecto al mínimo configurado.", 5, "MUST", "APROBADO", None, None),
                ("T0214", "Alertas de productos próximos a vencer (ventana configurable, ej. 7 días).", 5, "MUST", "APROBADO", None, None),
                ("T0215", "Movimientos de inventario: entrada, salida y ajuste con motivo y trazabilidad.", 5, "MUST", "APROBADO", None, None),
                ("T0216", "Órdenes de compra y recepción total o parcial de mercadería actualizando stock.", 8, "MUST", "APROBADO", None, None),
                ("T0217", "CRUD de proveedores con datos de contacto.", 5, "MUST", "APROBADO", None, None),
                ("T0218", "Registro de mermas con descuento automático de stock.", 5, "MUST", "APROBADO", None, None),
            ],
        ),
        (
            4,
            "HU-09",
            (
                "Como administrador o vendedor, quiero gestionar caja, gastos, devoluciones, crédito y clientes para cerrar el ciclo operativo diario."
            ),
            [
                ("T0219", "Apertura y cierre de caja con montos iniciales/finales y detección de diferencias.", 5, "MUST", "APROBADO", None, None),
                ("T0220", "Registro de gastos operativos por categoría.", 5, "MUST", "APROBADO", None, None),
                ("T0221", "Devoluciones de venta con restauración de stock y nota asociada.", 5, "MUST", "APROBADO", None, None),
                ("T0222", "Ventas al crédito: cuentas por cobrar y registro de pagos parciales o totales.", 8, "MUST", "APROBADO", None, None),
                ("T0223", "Registro de clientes y vinculación a ventas para seguimiento.", 5, "MUST", "APROBADO", None, None),
            ],
        ),
    ],
    "summary": [
        ("Sprint", "Sprint 2"),
        ("Período", "Semanas 5–9"),
        ("Fecha revisión", "Fin de semana 9"),
        ("Cumplimiento backlog", "100% (23 tareas / 4 HU)"),
        ("Horas aproximadas", "~100 h"),
        ("Incidente principal", "Ajustes integración SUNAT/OSE y pruebas en sandbox"),
        ("Feedback cliente destacado", "Operación diaria digital; siguiente foco: promociones, reportes e IA (Sprint 3)"),
    ],
}

SPRINT_3 = {
    "file": "SPRINT_REVIEW_3_CHILALO_SHOT.xlsx",
    "goal": (
        "Meta Sprint 3 (condensada): promociones en POS, pasarela de pagos (web), épica de fidelización por puntos/canje "
        "(historias por rol), app Android, reportes/dashboard, IA y hardware, pruebas y despliegue (semanas 10–14). "
        "Épica fidelización: «Como administrador, quiero premiar a clientes frecuentes con puntos y canjes» — desglosada en HU-12 a HU-15."
    ),
    "blocks": [
        (
            1,
            "HU-10",
            (
                "Como administrador, quiero configurar promociones y packs aplicables en el POS para aumentar ventas."
            ),
            [
                ("T0301", "Creación de packs de productos (ej. sixpack) y precios asociados.", 5, "MUST", "APROBADO", None, None),
                ("T0302", "Descuentos por porcentaje, monto fijo o volumen; reglas por fechas.", 5, "MUST", "APROBADO", None, None),
                ("T0303", "Promociones temporales y por categoría de producto.", 5, "MUST", "APROBADO", None, None),
                ("T0304", "Motor de aplicación automática de promociones al agregar ítems al carrito en POS web.", 8, "MUST", "APROBADO", None, None),
                ("T0305", "Pruebas de regresión POS con promociones activas.", 5, "MUST", "APROBADO", None, None),
            ],
        ),
        (
            2,
            "HU-11",
            (
                "Como sistema, quiero integrar una pasarela de pagos en el checkout web para cobrar con tarjeta/billetera de forma segura."
            ),
            [
                ("T0306", "Integración pasarela de pagos en ambiente sandbox (tokenización / intent de pago).", 8, "MUST", "APROBADO", None, None),
                ("T0307", "Confirmación de pago exitoso y actualización del estado de la venta en backend.", 5, "MUST", "APROBADO", None, None),
                ("T0308", "Manejo de rechazos, reintentos e idempotencia; logs para conciliación.", 5, "MUST", "APROBADO", None, None),
            ],
        ),
        (
            3,
            "HU-12",
            (
                "Como vendedor, quiero revisar las puntuaciones de los clientes según las ventas realizadas para atender y fidelizar."
            ),
            [
                ("T0309", "Consultar saldo y nivel de puntos del cliente desde POS o módulo de clientes.", 5, "MUST", "APROBADO", None, None),
                ("T0310", "Ver resumen de puntos generados por ventas asociadas al cliente.", 5, "MUST", "APROBADO", None, None),
                ("T0311", "Listado o ranking de clientes con mayor acumulación de puntos (filtros básicos).", 5, "SHOULD", "APROBADO", None, None),
            ],
        ),
        (
            4,
            "HU-13",
            (
                "Como administrador, quiero gestionar el ponderado de puntos y las reglas de premiación del programa de fidelización."
            ),
            [
                ("T0312", "Configurar equivalencia puntos/moneda y reglas base de acumulación por venta.", 5, "MUST", "APROBADO", None, None),
                ("T0313", "Ponderados o bonificaciones por categoría de producto o campaña (si aplica).", 5, "MUST", "APROBADO", None, None),
                ("T0314", "Panel de administración de parámetros del programa (activar/desactivar, límites).", 5, "MUST", "APROBADO", None, None),
                ("T0315", "Acumulación automática de puntos al confirmar venta según reglas vigentes.", 5, "MUST", "APROBADO", None, None),
            ],
        ),
        (
            5,
            "HU-14",
            (
                "Como cliente, quiero visualizar mi historial de puntos para conocer beneficios y promociones disponibles."
            ),
            [
                ("T0316", "Consulta de movimientos y saldo de puntos (web y/o app según alcance).", 5, "MUST", "APROBADO", None, None),
                ("T0317", "Visualización de promociones o beneficios alcanzables según puntos acumulados.", 5, "MUST", "APROBADO", None, None),
            ],
        ),
        (
            6,
            "HU-15",
            (
                "Como cliente, quiero canjear los puntos obtenidos por mis compras al pagar una nueva venta."
            ),
            [
                ("T0318", "Canje de puntos como descuento en checkout/POS con validación de saldo.", 5, "MUST", "APROBADO", None, None),
                ("T0319", "Registro auditable del canje y reglas mínimo/máximo de redención.", 5, "MUST", "APROBADO", None, None),
            ],
        ),
        (
            7,
            "HU-16",
            (
                "Como usuario móvil, quiero una app Android para consultar el negocio y vender con la misma lógica del backend."
            ),
            [
                ("T0320", "App Android: autenticación JWT y navegación principal (Kotlin + Jetpack Compose).", 5, "MUST", "APROBADO", None, None),
                ("T0321", "Dashboard móvil con KPIs (ventas del día, alertas de stock, indicadores clave).", 8, "MUST", "APROBADO", None, None),
                ("T0322", "POS móvil: búsqueda de productos, carrito, formas de pago y confirmación de venta.", 8, "MUST", "APROBADO", None, None),
                ("T0323", "Consulta de inventario, alertas y movimientos desde la app.", 5, "MUST", "APROBADO", None, None),
                ("T0324", "Historial de ventas con detalle y emisión/consulta de comprobante.", 5, "MUST", "APROBADO", None, None),
                ("T0325", "Sincronización estable con API REST (Retrofit) y manejo de errores de red.", 5, "MUST", "APROBADO", None, None),
            ],
        ),
        (
            8,
            "HU-17",
            (
                "Como propietario, quiero reportes y dashboard con exportación para analizar rentabilidad y decidir."
            ),
            [
                ("T0326", "Dashboard web: ventas día/semana/mes, comparativas y accesos rápidos.", 8, "MUST", "APROBADO", None, None),
                ("T0327", "Reportes de ventas con filtros (fecha, forma de pago, vendedor).", 5, "MUST", "APROBADO", None, None),
                ("T0328", "Reporte de inventario: valor de stock, productos con bajo giro.", 5, "MUST", "APROBADO", None, None),
                ("T0329", "Reporte de rentabilidad y gastos (visión de ganancias brutas/netas).", 5, "MUST", "APROBADO", None, None),
                ("T0330", "Exportación de reportes a PDF y Excel.", 5, "MUST", "APROBADO", None, None),
            ],
        ),
        (
            9,
            "HU-18",
            (
                "Como equipo, quiero IA de apoyo, hardware periférico, pruebas finales y despliegue en producción."
            ),
            [
                ("T0331", "Servicio de IA: recomendaciones de reabastecimiento según historial de ventas.", 8, "SHOULD", "APROBADO", None, "Python/FastAPI opcional"),
                ("T0332", "Predicción de demanda y análisis de rotación (productos fuertes/débiles).", 5, "SHOULD", "APROBADO", None, None),
                ("T0333", "Asistente/chatbot básico para consultas frecuentes del negocio (opcional según tiempo).", 5, "COULD", "APROBADO", None, None),
                ("T0334", "Integración lector de código de barras USB y soporte impresora térmica de tickets.", 5, "MUST", "APROBADO", None, None),
                ("T0335", "Pruebas funcionales e integración SUNAT en producción; corrección de hallazgos críticos.", 8, "MUST", "APROBADO", None, None),
                ("T0336", "Despliegue en nube, migración/carga de datos reales, capacitación y entrega de manual de usuario.", 8, "MUST", "APROBADO", None, None),
            ],
        ),
    ],
    "summary": [
        ("Sprint", "Sprint 3"),
        ("Período", "Semanas 10–14"),
        ("Fecha revisión", "Fin de semana 14"),
        ("Cumplimiento backlog", "100% (36 tareas en 9 historias/HU)"),
        ("Horas aproximadas", "~125 h"),
        ("Incidente principal", "Pasarela de pagos, IA/despliegue y pruebas en dispositivo Android"),
        ("Cierre", "Sistema listo para operación; documentación y capacitación entregadas"),
    ],
}

SPRINT_CONFIG = {1: SPRINT_1, 2: SPRINT_2, 3: SPRINT_3}


def unmerge_data_area(ws, start: int = 5, end: int = 80) -> None:
    for mr in list(ws.merged_cells.ranges):
        if mr.max_row < start or mr.min_row > end:
            continue
        try:
            ws.unmerge_cells(str(mr))
        except Exception:
            pass


def clear_data_rows(ws, start: int = 5, end: int = 80) -> None:
    for r in range(start, end + 1):
        for c in range(1, 11):
            ws.cell(row=r, column=c, value=None)


def merge_hu_columns(ws, r_start: int, r_end: int) -> None:
    if r_end > r_start:
        ws.merge_cells(f"A{r_start}:A{r_end}")
        ws.merge_cells(f"B{r_start}:B{r_end}")
        ws.merge_cells(f"C{r_start}:C{r_end}")


def write_blocks(ws, blocks, start_row: int = 5) -> int:
    r = start_row
    for orden, hu_id, hu_text, tasks in blocks:
        r_start = r
        first = tasks[0]
        tid, tdesc, est, prio, status, proc, obs = first
        ws.cell(row=r, column=1, value=orden)
        ws.cell(row=r, column=2, value=hu_id)
        ws.cell(row=r, column=3, value=hu_text)
        ws.cell(row=r, column=4, value=tid)
        ws.cell(row=r, column=5, value=tdesc)
        ws.cell(row=r, column=6, value=est)
        ws.cell(row=r, column=7, value=prio)
        ws.cell(row=r, column=8, value=status)
        ws.cell(row=r, column=9, value=proc)
        ws.cell(row=r, column=10, value=obs)
        r += 1
        for t in tasks[1:]:
            tid, tdesc, est, prio, status, proc, obs = t
            ws.cell(row=r, column=4, value=tid)
            ws.cell(row=r, column=5, value=tdesc)
            ws.cell(row=r, column=6, value=est)
            ws.cell(row=r, column=7, value=prio)
            ws.cell(row=r, column=8, value=status)
            ws.cell(row=r, column=9, value=proc)
            ws.cell(row=r, column=10, value=obs)
            r += 1
        merge_hu_columns(ws, r_start, r - 1)
    return r


def build(sprint_num: int) -> Path:
    cfg = SPRINT_CONFIG[sprint_num]
    blocks = cfg["blocks"]
    out_path = Path(__file__).resolve().parent / cfg["file"]

    if not TEMPLATE.exists():
        raise SystemExit(f"No se encontró la plantilla: {TEMPLATE}")

    wb = openpyxl.load_workbook(TEMPLATE)
    if SHEET_NAME not in wb.sheetnames:
        raise SystemExit(f"No se encontró la hoja '{SHEET_NAME}'. Hojas: {wb.sheetnames}")

    ws = wb[SHEET_NAME]
    ws["A1"] = "SPRINT GOAL"
    ws["D1"] = cfg["goal"]
    # Encabezado columna C: número de sprint en la tabla
    ws.cell(row=4, column=3, value=f"Historia de usuario (Sprint {sprint_num})")

    unmerge_data_area(ws, 5, 80)
    clear_data_rows(ws, 5, 80)
    next_row = write_blocks(ws, blocks, 5)
    last_data = next_row - 1

    if next_row < 69:
        ws.delete_rows(next_row, 69 - next_row)

    total_row = next_row
    ws.cell(row=total_row, column=4, value="TOTAL (estimación)")
    ws.cell(row=total_row, column=6, value=f"=SUM(F5:F{last_data})")

    base = total_row + 2
    for i, (k, v) in enumerate(cfg["summary"]):
        rr = base + i
        ws.cell(row=rr, column=1, value=k)
        ws.cell(row=rr, column=2, value=v)

    wb.save(out_path)
    print("Guardado:", out_path)
    return out_path


def main():
    ap = argparse.ArgumentParser(description="Genera Excel Sprint Review Chilalo Shot")
    ap.add_argument("sprint", type=int, choices=[1, 2, 3], help="Número de sprint (1, 2 o 3)")
    args = ap.parse_args()
    build(args.sprint)


if __name__ == "__main__":
    main()
