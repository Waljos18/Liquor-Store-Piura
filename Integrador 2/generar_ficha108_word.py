# -*- coding: utf-8 -*-
"""Genera FPIPS-108 Diseño del Sistema (Word) a partir del contenido de FICHA_08."""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


def set_cell_shading(cell, fill_hex: str):
    """fill_hex sin #, ej. D9E2F3"""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill_hex)
    tcPr.append(shd)


def add_table(doc, headers, rows, header_fill="D9E2F3"):
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.style = "Table Grid"
    hdr_cells = t.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].text = h
        set_cell_shading(hdr_cells[i], header_fill)
        for p in hdr_cells[i].paragraphs:
            for r in p.runs:
                r.bold = True
                r.font.size = Pt(10)
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = t.rows[ri + 1].cells[ci]
            cell.text = str(val) if val is not None else ""
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(10)
    doc.add_paragraph()


def mono_block(doc, text: str, size_pt=8):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.5)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    run.font.name = "Consolas"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
    run.font.size = Pt(size_pt)


def _strip_md_inline(s: str) -> str:
    return s.replace("**", "").replace("`", "")


def append_section6_from_markdown(doc, md_path: Path) -> None:
    """Inserta la sección 6 tal como está en FICHA_08_DISENO_SISTEMA.md (diagramas + tablas)."""
    if not md_path.is_file():
        doc.add_paragraph(f"(No se encontró el archivo: {md_path})")
        return
    text = md_path.read_text(encoding="utf-8")
    start = text.find("## 6. MODELO DE PERSISTENCIA")
    if start < 0:
        doc.add_paragraph("(No se encontró '## 6. MODELO DE PERSISTENCIA' en el markdown.)")
        return
    end = text.find("**Fecha de Actualización:", start)
    if end < 0:
        end = len(text)
    lines = text[start:end].splitlines()
    i = 0
    in_code = False
    code_buf: list[str] = []

    def is_table_sep(s: str) -> bool:
        s = s.strip()
        return s.startswith("|") and "---" in s

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if stripped.startswith("```"):
            if not in_code:
                in_code = True
                code_buf = []
            else:
                in_code = False
                body = "\n".join(code_buf).rstrip()
                if body:
                    mono_block(doc, body, 7)
                code_buf = []
            i += 1
            continue

        if in_code:
            code_buf.append(line)
            i += 1
            continue

        if stripped == "## 6. MODELO DE PERSISTENCIA":
            i += 1
            continue
        if stripped == "---" or stripped == "":
            i += 1
            continue

        if line.startswith("### ") and not line.startswith("#### "):
            doc.add_heading(_strip_md_inline(line[4:].strip()), level=2)
            i += 1
            continue
        if line.startswith("#### "):
            doc.add_heading(_strip_md_inline(line[5:].strip()), level=3)
            i += 1
            continue

        if stripped.startswith("|") and i + 1 < len(lines) and is_table_sep(lines[i + 1]):
            header_cells = [c.strip() for c in stripped.split("|")[1:-1]]
            i += 2
            rows: list[list[str]] = []
            while i < len(lines) and lines[i].strip().startswith("|") and not is_table_sep(lines[i]):
                row = [c.strip() for c in lines[i].split("|")[1:-1]]
                if len(row) == len(header_cells):
                    rows.append([_strip_md_inline(c) for c in row])
                i += 1
            hdr = [_strip_md_inline(h) for h in header_cells]
            if hdr and rows:
                add_table(doc, hdr, rows)
            continue

        if stripped:
            doc.add_paragraph(_strip_md_inline(stripped))
        i += 1


def main():
    out = Path(__file__).resolve().parent / "FPIPS-108_Diseño_Sistema_ChilaloShot.docx"
    doc = Document()

    # Portada
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Desarrollo de Sistemas de Información\n\n")
    r.bold = True
    r.font.size = Pt(14)
    r = p.add_run("Sistema Web y App Móvil para Gestión Integral de Licorería\n")
    r.bold = True
    r.font.size = Pt(13)
    r = p.add_run('Proyecto: Chilalo Shot\n\n')
    r.font.size = Pt(12)
    r = p.add_run("FPIPS-108 Diseño del Sistema de Información\n\n")
    r.bold = True
    r.font.size = Pt(12)
    r = p.add_run("Versión 1.0\n\nAbril 2025 – I\n\n")
    r.font.size = Pt(11)

    doc.add_paragraph()
    add_table(
        doc,
        ["N°", "Apellidos y Nombres"],
        [["1", "(Completar integrante 1)"], ["2", "(Completar integrante 2)"], ["3", "(Completar integrante 3)"]],
    )
    doc.add_page_break()

    # Índice (alineado con plantilla + secciones reales del contenido)
    doc.add_heading("Índice", level=1)
    idx = [
        "1. Historial del documento .................................................... 3",
        "2. Introducción .................................................................... 3",
        "3. Objetivos ......................................................................... 3",
        "4. Alcances .......................................................................... 3",
        "5. Diseño (ECU) y su propósito ...................................................... 3",
        "   5.1 Interfase móvil 01 — Autenticación / Login ................................. 3",
        "   5.2 Interfase 02 móvil — Dashboard ............................................. 4",
        "   5.3 Interfase 03 móvil — POS Android ........................................... 5",
        "   5.4 Interfase 04 web — POS React ............................................... 6",
        "   5.5 Interfase 05 web — Inventario .............................................. 7",
        "6. Modelo de persistencia ........................................................... 8",
        "   6.1 Diagrama entidad-relación — físico ....................................... 8",
        "   6.2 Relación Activity — ViewModel — Screen (Android) ......................... 9",
        "   6.3 Diagrama de clases ......................................................... 9",
        "   6.4 Diagrama de despliegue .................................................... 10",
        "   6.5 Diagrama de estados ........................................................ 10",
        "   6.6 Diagrama de secuencia ..................................................... 11",
        "   6.7 Diagrama de componentes ................................................... 11",
    ]
    for line in idx:
        doc.add_paragraph(line, style="List Bullet")

    doc.add_paragraph()
    p = doc.add_paragraph()
    p.add_run("Fecha de actualización: 05/04/2025    Versión: 1.0    Preparado por: Waljos18 / IDAT")
    doc.add_page_break()

    # 1 Historial
    doc.add_heading("1. Historial del documento", level=1)
    doc.add_heading("Información del documento", level=2)
    add_table(
        doc,
        ["Observaciones", "Modificado por", "Fecha"],
        [
            [
                "Creación inicial: diseño de interfaces móviles (Login, Dashboard, POS), web (POS, Ventas) y modelo de persistencia.",
                "Waljos18",
                "04/04/2025",
            ],
            [
                "Actualización: inclusión de módulos Inventario, Clientes, Ventas Android; diagramas UML completos.",
                "Waljos18",
                "05/04/2025",
            ],
        ],
    )

    # 2 Introducción
    doc.add_heading("2. Introducción", level=1)
    intro = (
        'El presente documento describe el diseño técnico del Sistema Web y Aplicación Móvil para la licorería '
        '"Chilalo Shot" ubicada en Piura, Perú. El proyecto comprende tres componentes principales:\n\n'
        "• Backend: API REST con Spring Boot 3.2.5 (Java 17), PostgreSQL 15 y seguridad JWT.\n"
        "• Frontend Web: React 18 + TypeScript + Vite con Tailwind CSS para administración y POS web.\n"
        "• Aplicación Android: Kotlin + Jetpack Compose + Hilt + Retrofit para gestión móvil en el punto de venta.\n\n"
        "El diseño sigue arquitectura en capas (Controller → Service → Repository → Entity) en el backend y MVVM "
        "(ViewModel → Repository → Remote/Local) en Android, con separación de responsabilidades, escalabilidad y mantenibilidad."
    )
    doc.add_paragraph(intro)

    # 3 Objetivos
    doc.add_heading("3. Objetivos", level=1)
    add_table(
        doc,
        ["N°", "Objetivo"],
        [
            ["1", "Definir la ECU para interfaces móviles y web del sistema Chilalo Shot."],
            ["2", "Documentar diseño visual y funcional de cada pantalla (componentes, herramientas, formularios)."],
            ["3", "Establecer modelo de persistencia (ER físico, clases, despliegue, secuencia y componentes)."],
            ["4", "Describir relación Activity–ViewModel–Screen en Android (Navigation + Compose)."],
            ["5", "Servir de referencia técnica para desarrollo e integración en el sprint de construcción."],
        ],
    )

    # 4 Alcances
    doc.add_heading("4. Alcances", level=1)
    doc.add_heading("4.1 Incluido", level=2)
    add_table(
        doc,
        ["Módulo", "Plataforma", "Descripción"],
        [
            ["Autenticación (Login/Logout)", "Android + Web", "JWT, redirección por rol"],
            ["Dashboard", "Android + Web", "KPIs del día, stock bajo, alertas"],
            ["Punto de Venta (POS)", "Android + Web", "Búsqueda, carrito, pagos, confirmar venta"],
            ["Inventario", "Android + Web", "Productos, stock, alertas bajo/vencimiento"],
            ["Clientes", "Android + Web", "CRUD, puntos de fidelización"],
            ["Ventas", "Android + Web", "Historial, filtros, detalle, comprobante"],
            ["Productos", "Android + Web", "Catálogo, categoría, búsqueda"],
            ["Facturación electrónica", "Web + Backend", "Boleta/factura SUNAT vía OSE"],
        ],
    )
    doc.add_heading("4.2 Excluido", level=2)
    for item in [
        "Módulo de contabilidad completa.",
        "Gestión de nómina de empleados.",
        "E-commerce / venta online.",
        "App para clientes finales.",
        "Módulo de delivery.",
    ]:
        doc.add_paragraph(item, style="List Bullet")

    # 5 ECU
    doc.add_heading("5. Diseño (ECU) y su propósito", level=1)

    doc.add_heading("5.1 Interfase móvil 01 — Autenticación / Login", level=2)
    add_table(
        doc,
        ["Cod. CU", "Nombre caso de uso", "RF", "Descripción"],
        [
            [
                "CU-MOB-01",
                "Iniciar sesión en app móvil",
                "RF-AUTH-01",
                "Validación vía POST /api/v1/auth/login; JWT en DataStore; navegación a MainScreen o error.",
            ],
            [
                "CU-MOB-02",
                "Cerrar sesión",
                "RF-AUTH-02",
                "Logout en TopAppBar; borra token y vuelve a LoginScreen.",
            ],
        ],
    )
    doc.add_heading("5.1.1 Diseño interfase móvil 01 — Login Screen", level=3)
    doc.add_paragraph("La pantalla muestra:")
    add_table(
        doc,
        ["ID", "Herramienta", "Formulario / descripción"],
        [
            ["IC-01", "Icon (Material3)", "Liquor 72 dp #2563EB"],
            ["LBL-01", "Text headlineMedium", "Chilalo Shot"],
            ["LBL-02", "Text bodyMedium", "Sistema de Gestión"],
            ["CARD-01", "ElevatedCard", "Formulario 16 dp"],
            ["LBL-03", "Text titleMedium", "Iniciar Sesión"],
            ["TXT-01", "OutlinedTextField", "Usuario"],
            ["TXT-02", "OutlinedTextField", "Contraseña + ojo"],
            ["LBL-ERR", "Text bodySmall rojo", "Error auth"],
            ["BTN-01", "Button filled", "Ingresar #2563EB"],
            ["PROG-01", "CircularProgressIndicator", "Carga"],
        ],
    )
    doc.add_paragraph(
        "Flujo: usuario y contraseña → BTN-01 → LoginViewModel → éxito MainScreen / error mensaje. "
        "Archivos: ui/login/LoginScreen.kt, LoginViewModel.kt"
    )

    doc.add_heading("5.2 Interfase 02 móvil — Dashboard / menú principal", level=2)
    add_table(
        doc,
        ["Cod. CU", "Nombre caso de uso", "RF", "Descripción"],
        [
            ["CU-MOB-03", "Ver dashboard", "RF-DASH-01", "KPIs vía GET /api/v1/reportes/dashboard."],
            ["CU-MOB-04", "Navegar módulos", "RF-NAV-01", "NavigationBar: Inicio, Caja, Productos, Ventas, Inventario, Clientes."],
            ["CU-MOB-05", "Configurar servidor", "RF-CFG-01", "ServerConfigDialog IP:puerto en TopAppBar."],
        ],
    )
    doc.add_paragraph("La pantalla muestra:")
    add_table(
        doc,
        ["ID", "Herramienta", "Formulario / descripción"],
        [
            ["TOPBAR-01", "CenterAlignedTopAppBar", "Título Chilalo Shot"],
            ["IC-WIFI", "IconButton", "Conectividad"],
            ["IC-REFRESH", "IconButton", "Recarga dashboard"],
            ["IC-LOGOUT", "IconButton", "Logout"],
            ["CARD-KPI-01..04", "ElevatedCard", "Ventas día, stock bajo, ingresos, productos activos"],
            ["NAV-01", "NavigationBar", "6 pestañas"],
            ["PROG-01", "CircularProgressIndicator", "Loading"],
            ["DIALOG-CFG", "ServerConfigDialog", "URL base"],
        ],
    )
    doc.add_paragraph("Archivos: DashboardScreen.kt, DashboardViewModel.kt, MainScreen.kt")

    doc.add_heading("5.3 Interfase 03 móvil — POS Android", level=2)
    add_table(
        doc,
        ["Cod. CU", "Nombre caso de uso", "RF", "Descripción"],
        [
            ["CU-MOB-06", "Buscar producto POS", "RF-POS-01", "GET /api/v1/productos?search=&categoriaId="],
            ["CU-MOB-07", "Agregar al carrito", "RF-POS-02", "Toque producto; badge carrito"],
            ["CU-MOB-08", "Gestionar carrito", "RF-POS-03", "ModalBottomSheet: qty, cliente, descuento, pago"],
            ["CU-MOB-09", "Confirmar venta", "RF-POS-04", "POST /api/v1/ventas"],
            ["CU-MOB-10", "Emitir comprobante", "RF-POS-05", "Boleta/factura POST facturación"],
        ],
    )
    doc.add_paragraph("La pantalla muestra: SearchBar, chips categoría, grid productos, FAB carrito, ModalBottomSheet.")
    add_table(
        doc,
        ["ID", "Herramienta", "Descripción breve"],
        [
            ["SEARCH-01", "SearchBar", "Búsqueda"],
            ["CHIPS-CAT", "LazyRow FilterChip", "Categorías"],
            ["GRID-PROD", "LazyVerticalGrid", "2 columnas"],
            ["FAB-CARRITO", "ExtendedFAB", "Ver carrito"],
            ["SHEET-01", "ModalBottomSheet", "Carrito y pago"],
            ["RADIO-PAGO", "RadioButton ×6", "Efectivo/Tarjeta/Yape/Plin/Transferencia/Mixto"],
            ["BTN-CONFIRM", "Button", "Confirmar venta"],
            ["DIALOG-OK", "AlertDialog", "Éxito boleta/factura"],
        ],
    )
    doc.add_paragraph("Archivos: POSScreen.kt, POSViewModel.kt, PosRepository.kt")

    doc.add_heading("5.4 Interfase 04 web — POS React", level=2)
    add_table(
        doc,
        ["Cod. CU", "Nombre caso de uso", "RF", "Descripción"],
        [
            ["CU-WEB-01", "Venta POS web", "RF-WEB-POS-01", "Búsqueda, carrito, MIXTO, boleta/factura"],
            ["CU-WEB-02", "Pago MIXTO", "RF-WEB-POS-02", "Dos métodos y montos"],
            ["CU-WEB-03", "Comprobante post-venta", "RF-WEB-POS-03", "Modales boleta/factura"],
        ],
    )
    add_table(
        doc,
        ["ID", "Herramienta", "Descripción"],
        [
            ["INPUT-SEARCH", "input Tailwind", "Debounce 300 ms"],
            ["CHIPS-CAT", "buttons", "Categorías"],
            ["TABLE-CARRITO", "table", "Qty inline, eliminar"],
            ["SELECT-PAGO", "select/radio", "Forma de pago"],
            ["PANEL-MIXTO", "div", "Solo si MIXTO"],
            ["BTN-COBRAR", "button", "POST /api/v1/ventas"],
        ],
    )
    doc.add_paragraph("Archivo: frontend/src/pages/POS.tsx")

    doc.add_heading("5.5 Interfase 05 web — Inventario", level=2)
    add_table(
        doc,
        ["Cod. CU", "Nombre caso de uso", "RF", "Descripción"],
        [
            ["CU-WEB-04", "Alertas inventario", "RF-INV-01", "Stock bajo y por vencer"],
            ["CU-WEB-05", "Movimiento manual", "RF-INV-02", "ENTRADA/SALIDA/AJUSTE + motivo"],
            ["CU-WEB-06", "Historial movimientos", "RF-INV-03", "Filtros y paginación servidor"],
        ],
    )

    # 6 Persistencia (contenido íntegro desde FICHA_08_DISENO_SISTEMA.md)
    doc.add_page_break()
    doc.add_heading("6. Modelo de persistencia", level=1)
    ficha_md = Path(__file__).resolve().parent / "FICHA_08_DISENO_SISTEMA.md"
    append_section6_from_markdown(doc, ficha_md)

    # Pie
    doc.add_paragraph()
    p = doc.add_paragraph()
    p.add_run("Fecha de actualización: 05/04/2025    Versión: 1.0    Preparado por: Waljos18    Documento: FPIPS-108")

    doc.save(out)
    print(f"Guardado: {out}")


if __name__ == "__main__":
    main()
