# -*- coding: utf-8 -*-
"""Genera FPIPS-109 Prueba de Calidad de Software (Word) desde FPIPS_109_PRUEBAS_CALIDAD_SOFTWARE.md."""
from pathlib import Path
import re
import shutil

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


def set_cell_shading(cell, fill_hex: str):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill_hex)
    tcPr.append(shd)


def add_table(doc, headers, rows, header_fill="D9E2F3"):
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.style = "Table Grid"
    hdr_cells = t.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].text = str(h)
        set_cell_shading(hdr_cells[i], header_fill)
        for p in hdr_cells[i].paragraphs:
            for r in p.runs:
                r.bold = True
                r.font.size = Pt(10)
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = t.rows[ri + 1].cells[ci]
            cell.text = str(val) if val is not None else ""
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(10)
    doc.add_paragraph()


def _strip_md(s: str) -> str:
    return s.replace("**", "").replace("`", "")


def extract_section(text: str, heading: str) -> str:
    marker = f"## {heading}"
    i = text.find(marker)
    if i < 0:
        return ""
    start = i + len(marker)
    j = text.find("\n## ", start)
    if j < 0:
        return text[start:].strip()
    return text[start:j].strip()


def parse_markdown_table_lines(lines: list[str], start: int) -> tuple[list[str], list[list[str]], int]:
    """Devuelve (headers, rows, next_index)."""
    if start >= len(lines) or not lines[start].strip().startswith("|"):
        return [], [], start
    hdr = [c.strip() for c in lines[start].split("|")[1:-1]]
    if start + 1 >= len(lines) or "---" not in lines[start + 1]:
        return [], [], start
    i = start + 2
    rows: list[list[str]] = []
    while i < len(lines) and lines[i].strip().startswith("|"):
        row = [c.strip() for c in lines[i].split("|")[1:-1]]
        if len(row) == len(hdr):
            rows.append([_strip_md(x) for x in row])
        i += 1
    return [_strip_md(h) for h in hdr], rows, i


def parse_all_tables(block: str) -> list[tuple[list[str], list[list[str]]]]:
    lines = block.splitlines()
    out: list[tuple[list[str], list[list[str]]]] = []
    i = 0
    while i < len(lines):
        if lines[i].strip().startswith("|"):
            h, r, ni = parse_markdown_table_lines(lines, i)
            if h and r:
                out.append((h, r))
                i = ni
            else:
                i += 1
        else:
            i += 1
    return out


def add_intro_paragraphs(doc, block: str):
    for line in block.splitlines():
        s = line.strip()
        if not s or s.startswith("#"):
            continue
        if s.startswith("- "):
            doc.add_paragraph(_strip_md(s[2:]), style="List Bullet")
        else:
            doc.add_paragraph(_strip_md(s))


def process_alcances(doc, block: str):
    lines = block.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()
        if stripped.startswith("### Fuera"):
            doc.add_heading("Fuera del alcance", level=2)
            i += 1
            while i < len(lines) and lines[i].strip().startswith("- "):
                doc.add_paragraph(_strip_md(lines[i].strip()[2:]), style="List Bullet")
                i += 1
            continue
        if stripped.startswith("### "):
            doc.add_heading(_strip_md(stripped[4:]), level=2)
            i += 1
            continue
        if line.strip().startswith("|"):
            h, r, ni = parse_markdown_table_lines(lines, i)
            if h and r:
                add_table(doc, h, r)
                i = ni
            else:
                i += 1
            continue
        if line.strip().startswith("- "):
            doc.add_paragraph(_strip_md(line.strip()[2:]), style="List Bullet")
            i += 1
            continue
        if line.strip():
            doc.add_paragraph(_strip_md(line.strip()))
        i += 1


def process_section8(doc, block: str):
    parts = re.split(r"(?=^### PRUEBA )", block, flags=re.MULTILINE)
    for part in parts:
        part = part.strip()
        if not part:
            continue
        lines = part.splitlines()
        if not lines[0].startswith("### PRUEBA"):
            continue
        title = _strip_md(lines[0].replace("### ", "").strip())
        doc.add_heading(title, level=2)
        rest = "\n".join(lines[1:]).strip()
        rest_lines = rest.splitlines()
        i = 0
        while i < len(rest_lines):
            if rest_lines[i].strip() == "---":
                i += 1
                continue
            if rest_lines[i].strip().startswith("#### "):
                doc.add_heading(_strip_md(rest_lines[i].strip()[5:]), level=3)
                i += 1
                continue
            if rest_lines[i].strip().startswith("|"):
                h, r, ni = parse_markdown_table_lines(rest_lines, i)
                if h and r:
                    add_table(doc, h, r)
                    i = ni
                else:
                    i += 1
                continue
            if rest_lines[i].strip():
                doc.add_paragraph(_strip_md(rest_lines[i].strip()))
            i += 1
        doc.add_paragraph()


def main():
    base = Path(__file__).resolve().parent
    md_path = base / "FPIPS_109_PRUEBAS_CALIDAD_SOFTWARE.md"
    out_local = base / "FPIPS-109_Prueba_Calidad_ChilaloShot.docx"
    out_downloads = Path(r"c:\Users\Usuario\Downloads\FPIPS-109 Prueba de Calidad de Software - Chilalo Shot.docx")

    text = md_path.read_text(encoding="utf-8")
    doc = Document()

    # --- Portada (desde cabecera MD hasta ---)
    pre = text.split("---", 1)[0]
    for line in pre.splitlines():
        s = line.strip()
        if not s or s.startswith("**Integrantes"):
            continue
        if s.startswith("# "):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(_strip_md(s[2:]))
            r.bold = True
            r.font.size = Pt(14)
        elif s.startswith("## "):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(_strip_md(s[3:]))
            r.bold = True
            r.font.size = Pt(12)
        elif s.startswith("### "):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(_strip_md(s[4:]))
            r.font.size = Pt(11)

    for h, rows in parse_all_tables(pre):
        if "Apellidos" in "".join(h):
            add_table(doc, h, rows)

    doc.add_page_break()

    # Índice
    doc.add_heading("Índice", level=1)
    for item in [
        "1. Historial del documento",
        "2. Introducción",
        "3. Objetivos",
        "4. Alcances",
        "5. Lista de requerimientos funcionales",
        "6. Lista de casos de uso",
        "7. Trazabilidad de pruebas — matriz RF vs CU",
        "8. Especificación de casos de prueba",
        "9. Resumen de resultados de pruebas",
    ]:
        doc.add_paragraph(item, style="List Bullet")
    doc.add_paragraph()
    doc.add_paragraph("Fecha de actualización: 05/04/2025    Versión: 1.0    Preparado por: Waljos18 / IDAT")
    doc.add_page_break()

    # §1 Historial
    doc.add_heading("1. Historial del documento", level=1)
    s1 = extract_section(text, "1. HISTORIAL DEL DOCUMENTO")
    for h, rows in parse_all_tables(s1):
        add_table(doc, h, rows)

    # §2 Introducción
    doc.add_heading("2. Introducción", level=1)
    s2 = extract_section(text, "2. INTRODUCCIÓN")
    add_intro_paragraphs(doc, s2)

    # §3 Objetivos
    doc.add_heading("3. Objetivos", level=1)
    s3 = extract_section(text, "3. OBJETIVOS")
    for h, rows in parse_all_tables(s3):
        add_table(doc, h, rows)

    # §4 Alcances
    doc.add_heading("4. Alcances", level=1)
    s4 = extract_section(text, "4. ALCANCES")
    process_alcances(doc, s4)

    # §5 RF
    doc.add_heading("5. Lista de requerimientos funcionales", level=1)
    s5 = extract_section(text, "5. LISTA DE REQUERIMIENTOS FUNCIONALES")
    for h, rows in parse_all_tables(s5):
        add_table(doc, h, rows)

    # §6 CU
    doc.add_heading("6. Lista de casos de uso", level=1)
    s6 = extract_section(text, "6. LISTA DE CASOS DE USO")
    for h, rows in parse_all_tables(s6):
        add_table(doc, h, rows)

    # §7 Matriz
    doc.add_heading("7. Trazabilidad de pruebas — matriz RF vs CU", level=1)
    s7 = extract_section(text, "7. TRAZABILIDAD DE PRUEBAS — MATRIZ RF vs CU")
    for h, rows in parse_all_tables(s7):
        add_table(doc, h, rows)

    # §8 Especificación
    doc.add_page_break()
    doc.add_heading("8. Especificación de casos de prueba", level=1)
    s8 = extract_section(text, "8. ESPECIFICACIÓN DE CASOS DE PRUEBA")
    process_section8(doc, s8)

    # §9 Resumen
    doc.add_page_break()
    doc.add_heading("9. Resumen de resultados de pruebas", level=1)
    s9 = extract_section(text, "9. RESUMEN DE RESULTADOS DE PRUEBAS")
    for h, rows in parse_all_tables(s9):
        add_table(doc, h, rows)
    for line in s9.splitlines():
        if line.strip().startswith("**Cobertura"):
            doc.add_paragraph(_strip_md(line.strip()))
        if line.strip().startswith("*Fecha de Actualización"):
            doc.add_paragraph(_strip_md(line.strip().lstrip("*")))

    doc.save(out_local)
    print(f"Guardado: {out_local}")
    try:
        shutil.copy2(out_local, out_downloads)
        print(f"Copiado a: {out_downloads}")
    except OSError as e:
        print(f"No se pudo copiar a Descargas ({e}). Usa el archivo en Integrador 2.")


if __name__ == "__main__":
    main()
