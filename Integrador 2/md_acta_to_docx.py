# -*- coding: utf-8 -*-
"""
Convierte ACTA_CONSTITUCION_PROYECTO.md a Word (.docx) con tablas y encabezados.
"""
import re
import shutil
from pathlib import Path

from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH


def add_paragraph_with_bold(doc, text, style=None):
    """Añade párrafo; **negrita** se convierte en runs."""
    p = doc.add_paragraph(style=style)
    parts = re.split(r"(\*\*.+?\*\*)", text)
    for part in parts:
        if part.startswith("**") and part.endswith("**"):
            run = p.add_run(part[2:-2])
            run.bold = True
        elif part:
            p.add_run(part)
    return p


def is_table_separator(line):
    return bool(re.match(r"^\|[\s\-:|]+\|\s*$", line.strip()))


def parse_table_row(line):
    line = line.strip()
    if not line.startswith("|"):
        return None
    cells = [c.strip() for c in line.strip("|").split("|")]
    return cells


def main():
    base = Path(__file__).resolve().parent
    md_path = base / "ACTA_CONSTITUCION_PROYECTO.md"
    # Destino principal en Documentos; si está abierto en Word, guardar copia en el proyecto
    out_documents = Path(r"c:\Users\Usuario\Documents\03 Acta de Constitución (1) (1).docx")
    out_fallback = base / "03_Acta_de_Constitucion_CORREGIDO.docx"
    out_path = out_documents

    if not md_path.exists():
        raise SystemExit(f"No se encontró: {md_path}")

    with open(md_path, encoding="utf-8") as f:
        lines = f.readlines()

    doc = Document()
    try:
        doc.styles["Normal"].font.name = "Calibri"
        doc.styles["Normal"].font.size = Pt(11)
    except Exception:
        pass

    i = 0
    in_table = False
    table_rows = []

    def flush_table():
        nonlocal table_rows, in_table
        if not table_rows:
            in_table = False
            return
        table_rows = [r for r in table_rows if any(c.strip() for c in r)]
        if not table_rows:
            in_table = False
            return
        ncols = max(len(r) for r in table_rows)
        for r in table_rows:
            while len(r) < ncols:
                r.append("")
        t = doc.add_table(rows=len(table_rows), cols=ncols)
        t.style = "Table Grid"
        for ri, row in enumerate(table_rows):
            for ci, cell in enumerate(row):
                t.rows[ri].cells[ci].text = cell
        doc.add_paragraph()
        table_rows = []
        in_table = False

    while i < len(lines):
        raw = lines[i]
        line = raw.rstrip("\n")
        stripped = line.strip()

        if stripped == "---":
            i += 1
            continue

        if stripped.startswith("|") and not is_table_separator(stripped):
            row = parse_table_row(stripped)
            if row:
                if not in_table:
                    in_table = True
                    table_rows = []
                table_rows.append(row)
            i += 1
            continue
        else:
            if in_table:
                flush_table()

        if stripped.startswith("# "):
            doc.add_heading(stripped[2:].strip(), level=0)
        elif stripped.startswith("## "):
            doc.add_heading(stripped[3:].strip(), level=1)
        elif stripped.startswith("### "):
            doc.add_heading(stripped[4:].strip(), level=2)
        elif stripped.startswith("#### "):
            doc.add_heading(stripped[5:].strip(), level=3)
        elif stripped.startswith("- "):
            add_paragraph_with_bold(doc, stripped[2:].strip(), style="List Bullet")
        elif stripped.startswith("> "):
            p = add_paragraph_with_bold(doc, stripped[2:].strip())
            for run in p.runs:
                run.italic = True
        elif re.match(r"^\d+\.\s+", stripped):
            body = re.sub(r"^\d+\.\s*", "", stripped)
            add_paragraph_with_bold(doc, body, style="List Number")
        elif stripped == "":
            pass
        else:
            add_paragraph_with_bold(doc, stripped)

        i += 1

    if in_table:
        flush_table()

    # Pie de documento (líneas finales italic en md)
    doc.add_paragraph()
    p = doc.add_paragraph(
        "Documento elaborado como parte del Trabajo Académico Aplicado (TAA). "
        "Instituto de Educación Superior – Piura, Perú. Versión 1.0 – Marzo 2025."
    )
    for run in p.runs:
        run.italic = True

    try:
        if out_path.exists():
            shutil.copy2(out_path, out_path.with_suffix(".docx.bak"))
        doc.save(str(out_path))
        print("Guardado:", out_path)
    except OSError as e:
        doc.save(str(out_fallback))
        print("No se pudo escribir en Documentos (¿archivo abierto en Word?).")
        print("Guardado en proyecto:", out_fallback)
        print("Detalle:", e)


if __name__ == "__main__":
    main()
