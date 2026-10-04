# -*- coding: utf-8 -*-
"""
Genera Introducción + Capítulo I (Análisis del negocio) del TAAP en Word (.docx),
mismo criterio de formato que generar_cap_iv_taap.py / generar_cap_v_taap.py.
Contexto: Licorería Chilalo Shot, Piura — sistema integral POS/ERP.
"""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_LINE_SPACING, WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt

OUT_FILE = Path(__file__).resolve().parent / "CAPITULO_I_INTRODUCCION_TAAP_APA7.docx"

KEEP_UPPER = {
    "I", "II", "III", "IV", "V", "API", "JWT", "SQL", "PUCP", "USMP", "USIL", "TAAP",
    "SUNAT", "BPMN", "POS", "ERP", "RF", "RNF", "IGV", "OSE", "TI", "TIC",
}


def set_run_font(run, name: str = "Times New Roman", size: int = 12, bold: bool = False, italic: bool = False) -> None:
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic


def configure_section(section) -> None:
    section.page_height = Inches(11)
    section.page_width = Inches(8.5)
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(2.54)
    section.right_margin = Cm(2.54)


def strip_heading_number(s: str) -> str:
    return re.sub(r"^\d+(\.\d+)*\s*", "", s).strip()


def apa_title_case(s: str) -> str:
    s = strip_heading_number(s)
    roman = {"i": "I", "ii": "II", "iii": "III", "iv": "IV", "v": "V"}

    def one_word(raw: str) -> str:
        trailing = ""
        if raw.endswith(":"):
            raw, trailing = raw[:-1], ":"
        core = re.sub(r"^\W+|\W+$", "", raw)
        if not core:
            return raw + trailing
        low = core.lower()
        if low in roman:
            return roman[low] + trailing
        if core.upper() in KEEP_UPPER:
            return core.upper() + trailing
        return core[:1].upper() + core[1:].lower() + trailing

    return " ".join(one_word(w) for w in s.split())


def add_heading_centered(doc: Document, text: str, size: int = 14) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.first_line_indent = Cm(0)
    run = p.add_run(text)
    set_run_font(run, size=size, bold=True)


def add_heading_h2(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.first_line_indent = Cm(0)
    run = p.add_run(text)
    set_run_font(run, size=14, bold=True)


def add_heading_h3(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.first_line_indent = Cm(0)
    run = p.add_run(text)
    set_run_font(run, size=12, bold=True, italic=True)


def add_normal_paragraph(doc: Document, text: str) -> None:
    text = text.strip()
    if not text:
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.first_line_indent = Inches(0.5)
    p.paragraph_format.widow_control = True
    run = p.add_run(text)
    set_run_font(run)


def add_figure_caption(doc: Document, n: int, title: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    p.paragraph_format.first_line_indent = Cm(0)
    run = p.add_run(f"Figura {n}. {title}")
    set_run_font(run, size=12, bold=True)


def add_placeholder_block(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    set_run_font(run, size=11, italic=True)


def add_bullet_list(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
        p.paragraph_format.left_indent = Inches(0.5)
        p.paragraph_format.first_line_indent = Inches(-0.25)
        run = p.add_run("• " + item)
        set_run_font(run)


def add_table(
    doc: Document,
    caption: str | None,
    headers: list[str],
    rows: list[list[str]],
) -> None:
    if caption:
        p = doc.add_paragraph()
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
        p.paragraph_format.first_line_indent = Cm(0)
        run = p.add_run(caption)
        set_run_font(run, size=12, bold=True)
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    for j, h in enumerate(headers):
        cell = table.rows[0].cells[j]
        cell.text = ""
        cp = cell.paragraphs[0]
        cp.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        run = cp.add_run(h)
        set_run_font(run, size=12, bold=True)
    for i, row in enumerate(rows, start=1):
        for j in range(len(headers)):
            txt = row[j] if j < len(row) else ""
            cell = table.rows[i].cells[j]
            cell.text = ""
            cp = cell.paragraphs[0]
            cp.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
            run = cp.add_run(txt)
            set_run_font(run, size=12)
    doc.add_paragraph()


def build_document() -> Document:
    doc = Document()
    for section in doc.sections:
        configure_section(section)

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT

    # ---------- INTRODUCCIÓN ----------
    add_heading_centered(doc, "Introducción")
    add_normal_paragraph(
        doc,
        "El presente documento describe el trabajo aplicado para el diseño y la implementación de un "
        "sistema integral de gestión orientado a la licorería Chilalo Shot, ubicada en la ciudad de Piura. "
        "La investigación aplicada se centra en cómo las tecnologías de información pueden mejorar procesos "
        "de punto de venta, inventario, compras, facturación electrónica y atención al cliente en un "
        "comercio minorista de bebidas alcohólicas y productos afines, en un contexto de creciente "
        "exigencia normativa y de digitalización.",
    )
    add_normal_paragraph(
        doc,
        "La motivación del trabajo radica en reducir errores operativos, acelerar el registro de ventas, "
        "fortalecer el control de stock y disponer de información oportuna para la toma de decisiones, "
        "sin descuidar la trazabilidad frente a SUNAT mediante comprobantes electrónicos. El enfoque "
        "combina análisis del negocio, ingeniería de requisitos, arquitectura de software en tres capas y "
        "validación práctica mediante pruebas sobre la API y los clientes desarrollados.",
    )
    add_heading_h3(doc, "Objetivos")
    add_bullet_list(
        doc,
        [
            "Objetivo general: Proponer e implementar una solución de software (backend, aplicación web y "
            "aplicación móvil) que automatice los procesos críticos de venta, inventario y reportes de la "
            "organización cliente, alineada a la normativa vigente y a las buenas prácticas de seguridad.",
            "Objetivo específico 1: Caracterizar el negocio, identificar necesidades y documentar requisitos "
            "funcionales y no funcionales del sistema.",
            "Objetivo específico 2: Definir la arquitectura, el modelo de datos y los módulos principales "
            "del producto software.",
            "Objetivo específico 3: Implementar y probar la solución, documentando la programación, las "
            "pruebas de calidad y las conclusiones del proyecto.",
        ],
    )
    add_normal_paragraph(
        doc,
        "La metodología se apoyó en revisión documental, entrevistas y observación de procesos (elicitation), "
        "modelado de requisitos, desarrollo iterativo sobre stack Java/Spring Boot, React y Android "
        "(Kotlin), y pruebas integrales sobre servicios REST. El contenido del informe se organiza para que "
        "el lector comprenda primero el negocio y la necesidad (Capítulo I), seguido del marco técnico y "
        "funcional en capítulos posteriores, la implementación (Capítulo IV), las pruebas y cierre (Capítulo V), "
        "y los anexos de respaldo.",
    )
    add_heading_h3(doc, "Estructura Del Documento")
    add_bullet_list(
        doc,
        [
            "Capítulo I. Análisis del negocio: organización, FODA, procesos, necesidades, requisitos, "
            "propuesta y factibilidad.",
            "Capítulos intermedios (según versión final del TAAP): marco teórico, requisitos detallados o "
            "diseño, según lo establecido por la institución.",
            "Capítulo IV. Programación: arquitectura, base de datos, dependencias y codificación del backend, "
            "frontend y reportes.",
            "Capítulo V. Pruebas de calidad de software, conclusiones, recomendaciones y referencias; "
            "anexos con evidencias extensas.",
        ],
    )

    doc.add_page_break()

    # ---------- CAPÍTULO I ----------
    add_heading_centered(doc, "Capítulo I. Análisis Del Negocio")

    add_heading_h2(doc, "Generalidades")
    add_normal_paragraph(
        doc,
        "Este capítulo presenta de manera sintética la realidad de la organización cliente, sus procesos "
        "relevantes y la justificación del proyecto de sistemas que da origen al producto software desarrollado. "
        "La información se orienta a sustentar decisiones de alcance, requisitos y factibilidad.",
    )

    add_heading_h2(doc, "Descripción De La Organización")
    add_normal_paragraph(
        doc,
        "La licorería Chilalo Shot es una organización dedicada a la venta al por menor de licores, vinos, "
        "cervezas y productos complementarios, con atención en local y cobertura principal en el mercado "
        "piurano. Opera con personal de mostrador, almacén y administración, gestiona proveedores recurrentes "
        "y debe cumplir obligaciones tributarias y de registro de ventas ante SUNAT. El proyecto consiste en "
        "diseñar y ejecutar un sistema denominado en adelante sistema integral de gestión para licorería, que "
        "centralice catálogo, stock, ventas, compras, caja, reportes y facturación electrónica en la medida "
        "de lo definido en el alcance.",
    )
    add_normal_paragraph(
        doc,
        "NOTA: Si la institución docente exige trabajar con una empresa estrictamente ficticia, puede "
        "sustituirse el nombre comercial anterior por una denominación genérica y añadirse la siguiente "
        "frase al documento final: «La organización citada es un caso elaborado para fines académicos, "
        "ajustado a un perfil típico de PYME comercial en el Perú.»",
    )

    add_heading_h2(doc, "Historia De La Organización")
    add_normal_paragraph(
        doc,
        "A continuación se resume una línea de tiempo de hitos representativos (valores orientativos para "
        "fines académicos; deben contrastarse con datos reales si el trabajo se vincula a una empresa existente):",
    )
    add_bullet_list(
        doc,
        [
            "Año fundacional: constitución del negocio familiar de abarrotes y bebidas con foco en barrio.",
            "Expansión de surtido: incorporación de licores importados y nacionales; crecimiento de clientes "
            "mayoristas menores.",
            "Formalización tributaria: emisión de comprobantes y afiliación a operador de servicios electrónicos.",
            "Situación actual: mayor volumen de transacciones, necesidad de control de inventario en tiempo "
            "casi real y de canales de pago digitales (Yape, Plin, mixtos).",
        ],
    )

    add_heading_h2(doc, "Visión Empresarial")
    add_normal_paragraph(
        doc,
        "«Ser en Piura una licorería de referencia por la variedad de su oferta, la confianza en la atención "
        "y la eficiencia en la gestión, integrando canales presenciales y digitales que mejoren la experiencia "
        "del cliente y la rentabilidad sostenible del negocio.»",
    )

    add_heading_h2(doc, "Misión Empresarial")
    add_normal_paragraph(
        doc,
        "«Comercializar productos de licorería con estándares de calidad y precios competitivos, asegurando "
        "un servicio ágil en punto de venta, cumplimiento normativo en facturación y un manejo responsable "
        "del inventario, en beneficio de clientes, colaboradores y proveedores.»",
    )

    add_heading_h2(doc, "Organigrama Y Funciones De Las Principales Áreas")
    add_normal_paragraph(
        doc,
        "Se presenta un organigrama simplificado de hasta tres niveles jerárquicos, mostrando solo las áreas "
        "principales. El área donde se ejecuta el proyecto —Operaciones y punto de venta / almacén— se destaca "
        "porque concentra el registro de ventas, el despacho y la interacción con inventario.",
    )
    add_figure_caption(doc, 1, "Organigrama Simplificado De Chilalo Shot (Tres Niveles)")
    add_placeholder_block(
        doc,
        "[Insertar figura: Gerencia → Administración y Finanzas; Operaciones (Punto de venta y Almacén); "
        "Logística/Compras. Resaltar el recuadro de Operaciones como ámbito principal del proyecto.]",
    )
    add_normal_paragraph(
        doc,
        "Gerencia define políticas comerciales y aprueba inversiones en TI. Administración y finanzas supervisa "
        "caja, conciliaciones y reportes. Operaciones ejecuta ventas, control de stock en piso y atención al "
        "cliente; es el principal usuario del POS y de la aplicación móvil. Compras gestiona proveedores y "
        "recepción de mercadería, vinculada al módulo de compras del sistema.",
    )

    add_heading_h2(doc, "Análisis FODA")
    add_normal_paragraph(
        doc,
        "La matriz FODA resume factores internos y externos. Se incluyen tres ítems por cuadrante, en el "
        "rango solicitado.",
    )
    add_table(
        doc,
        "Tabla 1. Análisis FODA De La Organización Cliente",
        ["Dimensión", "Descripción"],
        [
            [
                "Fortaleza 1",
                "Equipo conocedor del catálogo y de la clientela frecuente; flexibilidad para promociones.",
            ],
            [
                "Fortaleza 2",
                "Ubicación y relación con proveedores que permiten reposición relativamente rápida.",
            ],
            [
                "Fortaleza 3",
                "Disposición a adoptar medios de pago digitales y canales de comunicación con clientes.",
            ],
            [
                "Debilidad 1",
                "Registro manual o semiautomatizado de inventario, con riesgo de descuadres y mermas no "
                "trazadas.",
            ],
            [
                "Debilidad 2",
                "Dependencia de pocas personas para operar el POS y elaborar reportes para gerencia.",
            ],
            [
                "Debilidad 3",
                "Carga administrativa asociada a SUNAT y a la consolidación de ventas en hojas de cálculo.",
            ],
            [
                "Oportunidad 1",
                "Demanda sostenida de consumo organizado y entrega a domicilio asociada a eventos locales.",
            ],
            [
                "Oportunidad 2",
                "Programas de fidelización y venta cruzada mediante datos de compra histórica.",
            ],
            [
                "Oportunidad 3",
                "Financiamiento o asesoría para digitalización de MYPE en la región.",
            ],
            [
                "Amenaza 1",
                "Competencia de cadenas y licorerías con mayor automatización y precios agresivos.",
            ],
            [
                "Amenaza 2",
                "Riesgos de incumplimiento o demora en facturación electrónica ante cambios normativos.",
            ],
            [
                "Amenaza 3",
                "Variabilidad del tipo de cambio y costos de importación que comprimen margen.",
            ],
        ],
    )

    add_heading_h2(doc, "Mapa De Procesos")
    add_normal_paragraph(
        doc,
        "El mapa de procesos ofrece una visión de extremo a extremo. Se distinguen procesos de gestión "
        "(dirección y planeamiento), procesos primarios (misionales: venta, compra, atención al cliente) y "
        "procesos de soporte (recursos humanos básicos, TI, mantenimiento y cumplimiento tributario simplificado).",
    )
    add_figure_caption(doc, 2, "Mapa De Procesos (Gestión, Primarios Y Soporte)")
    add_placeholder_block(
        doc,
        "[Insertar lámina: bloque superior Gestión estratégica; bloque central Ventas, Compras, Inventario, "
        "Atención; bloque inferior Soporte: TI, Administración, Cumplimiento. Utilizar notación acordada por "
        "el curso.]",
    )

    add_heading_h2(doc, "Proceso(S) A Automatizar (Modelo Actual As-Is En BPMN)")
    add_normal_paragraph(
        doc,
        "El foco de automatización es el proceso de venta en mostrador (y su extensión a escenarios con "
        "pago mixto y crédito de cliente), incluyendo validación de stock, emisión o preparación de "
        "comprobante y actualización de inventario. El modelo As-Is describe el flujo actual antes de la "
        "solución: recepción del pedido, búsqueda de precio, registro en sistema o papel, cobro y entrega.",
    )
    add_figure_caption(doc, 3, "Diagrama BPMN As-Is Del Proceso De Venta En Mostrador")
    add_placeholder_block(
        doc,
        "[Insertar modelo BPMN 2.0: inicio, tareas de atención, decisión de stock, registro de venta, cobro, "
        "fin; swimlanes Vendedor y Caja si aplica. Exportar desde Camunda Modeler, Bizagi o herramienta equivalente.]",
    )

    add_heading_h2(doc, "Identificación De Las Necesidades (Problema U Oportunidad)")
    add_normal_paragraph(
        doc,
        "La organización requiere superar limitaciones asociadas a la dispersión de la información entre "
        "mostrador, almacén y administración. Los principales dolores identificados son demoras en cuadre de "
        "caja, dificultad para anticipar quiebres de stock, reprocesos al consolidar ventas del día y exposición "
        "a errores en precios o promociones. Asimismo, existe la oportunidad de mejorar la experiencia del "
        "cliente mediante tiempos de atención más breves y medios de pago alineados al mercado local.",
    )
    add_normal_paragraph(
        doc,
        "Estas situaciones pueden atenuarse con un sistema que unifique catálogo y precios, descuente stock "
        "en cada venta, registre formas de pago detalladas y genere reportes y comprobantes de manera "
        "estructurada, con roles diferenciados para administrador y vendedor.",
    )

    add_heading_h2(doc, "Elicitación De Requisitos")
    add_normal_paragraph(
        doc,
        "A partir de la necesidad anterior se priorizan requisitos funcionales (RF) y no funcionales (RNF). "
        "La lista es orientativa y puede ampliarse en anexos.",
    )
    add_table(
        doc,
        "Tabla 2. Requisitos Prioritarios",
        ["Id.", "Tipo", "Descripción"],
        [
            ["RF-01", "RF", "Autenticación de usuarios con roles administrador y vendedor."],
            ["RF-02", "RF", "Registro de venta con carrito, descuentos y formas de pago (efectivo, tarjeta, Yape, Plin, mixto, crédito)."],
            ["RF-03", "RF", "Gestión de productos, categorías, stock y alertas de mínimo."],
            ["RF-04", "RF", "Gestión de clientes y puntos de fidelización."],
            ["RF-05", "RF", "Compras a proveedores y recepción de mercadería."],
            ["RF-06", "RF", "Reportes de ventas, inventario y dashboard operativo."],
            ["RF-07", "RF", "Integración o preparación para comprobantes electrónicos SUNAT (según alcance)."],
            ["RF-08", "RF", "Cierre de caja y registro de gastos/mermas básicos."],
            ["RNF-01", "RNF", "Tiempo de respuesta aceptable en operaciones frecuentes de venta (< 3 s objetivo en red local)."],
            ["RNF-02", "RNF", "Seguridad: contraseñas cifradas, API con JWT y HTTPS en despliegue."],
            ["RNF-03", "RNF", "Disponibilidad razonable en horario comercial; respaldos de base de datos."],
            ["RNF-04", "RNF", "Usabilidad: interfaces claras en web y móvil para usuarios con formación básica."],
            ["RNF-05", "RNF", "Mantenibilidad: código modular y migraciones versionadas de esquema (Flyway)."],
        ],
    )

    add_heading_h2(doc, "Análisis (Del Problema U Oportunidad)")
    add_normal_paragraph(
        doc,
        "El análisis busca precisar la naturaleza del problema, sus causas y el impacto en la organización.",
    )
    add_heading_h3(doc, "Definición Del Problema U Oportunidad")
    add_normal_paragraph(
        doc,
        "El problema central es la insuficiente integración de datos de venta e inventario en tiempo operativo; "
        "la oportunidad es convertir esos datos en decisiones de surtido, precios y fidelización.",
    )
    add_heading_h3(doc, "Causas Raíz Y Factores Limitantes")
    add_normal_paragraph(
        doc,
        "Entre las causas figuran herramientas no integradas, rotación de personal sin protocolos estandarizados "
        "y carga manual de reportes. Factores externos incluyen presión competitiva y normativa fiscal.",
    )
    add_heading_h3(doc, "Impacto En La Organización")
    add_normal_paragraph(
        doc,
        "El impacto se manifiesta en horas-hombre dedicadas a cuadres, posible pérdida de ventas por falta de "
        "stock visible y riesgos de desalineación entre precio mostrador y sistema administrativo.",
    )
    add_heading_h3(doc, "Stakeholders")
    add_normal_paragraph(
        doc,
        "Se identifican como partes interesadas: propietarios/gerencia, vendedores, personal de almacén, "
        "administración, proveedores recurrentes y clientes habituales; cada uno con expectativas de "
        "confiabilidad, rapidez y cumplimiento.",
    )

    add_heading_h2(doc, "Propuesta De Solución")
    add_normal_paragraph(
        doc,
        "La solución propuesta atiende tres perspectivas. En procesos, automatiza venta, inventario, compras, "
        "reportes y trazabilidad de caja. En personas, capacita roles diferenciados y reduce tareas repetitivas. "
        "En infraestructura tecnológica, contempla servidores o equipos para backend y base PostgreSQL, "
        "estaciones con navegador para el panel web, dispositivos Android para vendedores, y servicios TIC "
        "(conectividad, respaldo, certificados en producción).",
    )

    add_heading_h2(doc, "Alternativas A La Solución Propuesta")
    add_normal_paragraph(
        doc,
        "Se comparan, a nivel ejecutivo, al menos dos alternativas además de la desarrollada en el proyecto. "
        "Los costos son estimaciones orientativas en soles (PEN) y deben actualizarse con cotizaciones reales.",
    )
    add_table(
        doc,
        "Tabla 3. Alternativas De Solución (Ejecutivo)",
        ["Alternativa", "Descripción", "Costo Aproximado De Adquisición"],
        [
            [
                "Adquisición de software POS comercial cerrado",
                "Licencia por puesto o suscripción mensual con soporte; rápida puesta en marcha pero menor "
                "flexibilidad de personalización y dependencia del proveedor.",
                "Desde aprox. S/ 1500 – S/ 8 000 iniciales + S/ 100 – S/ 400 mensuales (según marca y módulos).",
            ],
            [
                "Desarrollo a medida por terceros",
                "Contrato con consultora para entregar sistema similar; reduce riesgo técnico interno pero "
                "incrementa costo y plazos de especificación.",
                "Desde aprox. S/ 15 000 – S/ 45 000 según alcance y integraciones SUNAT.",
            ],
            [
                "Solución propuesta (proyecto académico / desarrollo propio)",
                "Stack Spring Boot, React y Android bajo control del negocio; costo principalmente tiempo, "
                "infraestructura y mantenimiento.",
                "Infraestructura y despliegue estimados S/ 800 – S/ 3 500 primer año (hardware básico, hosting, "
                "dominio); horas de desarrollo no monetizadas en contexto académico.",
            ],
        ],
    )

    add_heading_h2(doc, "Factibilidad Del Proyecto")
    add_heading_h3(doc, "Recursos Tecnológicos Y Factibilidad Técnica")
    add_normal_paragraph(
        doc,
        "La organización puede operar el sistema con computadoras actuales, conexión a internet estable para "
        "actualizaciones y facturación electrónica, y dispositivos móviles para vendedores. Riesgos tecnológicos "
        "incluyen interrupciones de red, obsolescencia de equipos y la curva de aprendizaje de despliegue "
        "(Docker, servidor, certificados). Mitigación: redundancia básica, respaldos y documentación.",
    )
    add_heading_h3(doc, "Recursos Humanos Y Factibilidad Operativa")
    add_normal_paragraph(
        doc,
        "El personal de mostrador puede asimilar interfaces sencillas tras capacitación breve; la estructura "
        "actual es compatible si se asigna un responsable de TI interno o un proveedor de soporte. Los procesos "
        "impactados son venta, inventario, compras y reportes; deben mapearse responsables por tarea.",
    )
    add_heading_h3(doc, "Análisis Costo-Beneficio Y Factibilidad Económica")
    add_normal_paragraph(
        doc,
        "Los costos incluyen hardware, eventual hosting, tiempo de capacitación y mantenimiento. Los beneficios "
        "tangibles incluyen reducción de errores de cobro y mejores decisiones de compra; los intangibles, "
        "imagen profesional y trazabilidad tributaria. La relación costo-beneficio es favorable cuando se "
        "valoran las horas ahorradas en cuadre y la disminución de faltantes no detectados.",
    )
    add_normal_paragraph(
        doc,
        "En conclusión, la factibilidad del proyecto es un análisis integral que permite determinar si es posible, "
        "viable y rentable llevarlo a cabo. Con los recursos descritos y el alcance acotado a una PYME licorera, "
        "el desarrollo del sistema integral es técnicamente abordable, operativamente asumible con capacitación "
        "y económicamente razonable frente a alternativas de licencia cerrada, sirviendo como insumo para la "
        "toma de decisiones estratégicas.",
    )

    return doc


def main() -> None:
    doc = build_document()
    doc.save(OUT_FILE)
    print(f"Guardado: {OUT_FILE}")


if __name__ == "__main__":
    main()
