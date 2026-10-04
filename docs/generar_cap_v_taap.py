# -*- coding: utf-8 -*-
"""
Genera el Capítulo V del TAAP (Pruebas de calidad, conclusiones, referencias, anexos)
en Word (.docx) con el mismo criterio de formato que generar_cap_iv_taap.py (APA 7 / TAAP).

Las pruebas integrales se documentan a partir de la colección Postman
`LICORERIA_BACKEND.postman_collection.json` (misma carpeta que este script).
"""
from __future__ import annotations

import json
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_LINE_SPACING, WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt

OUT_FILE = Path(__file__).resolve().parent / "CAPITULO_V_TAAP_APA7.docx"
POSTMAN_COLLECTION = Path(__file__).resolve().parent / "LICORERIA_BACKEND.postman_collection.json"

KEEP_UPPER = {"IV", "V", "III", "II", "API", "JWT", "SQL", "PUCP", "USMP", "USIL", "TAAP"}


def set_run_font(run, name: str = "Times New Roman", size: int = 12, bold: bool = False, italic: bool = False) -> None:
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic


def configure_section(section) -> None:
    section.page_height = Inches(11)
    section.page_width = Inches(8.5)
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(2.54)
    section.right_margin = Cm(2.54)


def strip_heading_number(s: str) -> str:
    return re.sub(r"^\d+(\.\d+)*\s*", "", s).strip()


def apa_title_case(s: str) -> str:
    s = strip_heading_number(s)
    roman = {"ii": "II", "iii": "III", "iv": "IV", "v": "V"}

    def one_word(raw: str) -> str:
        trailing = ""
        if raw.endswith(":"):
            raw, trailing = raw[:-1], ":"
        core = re.sub(r"^\W+|\W+$", "", raw)
        if not core:
            return raw + trailing
        low = core.lower()
        if low in roman:
            return roman[low] + trailing
        if core.upper() in KEEP_UPPER:
            return core.upper() + trailing
        return core[:1].upper() + core[1:].lower() + trailing

    return " ".join(one_word(w) for w in s.split())


def add_heading_centered(doc: Document, text: str, size: int = 14) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.first_line_indent = Cm(0)
    run = p.add_run(text)
    set_run_font(run, size=size, bold=True)


def add_heading_h2(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.first_line_indent = Cm(0)
    run = p.add_run(text)
    set_run_font(run, size=14, bold=True)


def add_heading_h3(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.first_line_indent = Cm(0)
    run = p.add_run(text)
    set_run_font(run, size=12, bold=True, italic=True)


def add_normal_paragraph(doc: Document, text: str) -> None:
    text = text.strip()
    if not text:
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.first_line_indent = Inches(0.5)
    p.paragraph_format.widow_control = True
    run = p.add_run(text)
    set_run_font(run)


def add_bullet_list(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
        p.paragraph_format.left_indent = Inches(0.5)
        p.paragraph_format.first_line_indent = Inches(-0.25)
        run = p.add_run("• " + item)
        set_run_font(run)


def add_ref_entry(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    p.paragraph_format.first_line_indent = Cm(-1.27)
    p.paragraph_format.left_indent = Cm(1.27)
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(text)
    set_run_font(run, size=12)


def add_table(
    doc: Document,
    caption: str | None,
    headers: list[str],
    rows: list[list[str]],
) -> None:
    if caption:
        p = doc.add_paragraph()
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
        p.paragraph_format.first_line_indent = Cm(0)
        run = p.add_run(caption)
        set_run_font(run, size=12, bold=True)
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    for j, h in enumerate(headers):
        cell = table.rows[0].cells[j]
        cell.text = ""
        cp = cell.paragraphs[0]
        cp.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        run = cp.add_run(h)
        set_run_font(run, size=11, bold=True)
    for i, row in enumerate(rows, start=1):
        for j in range(len(headers)):
            txt = row[j] if j < len(row) else ""
            cell = table.rows[i].cells[j]
            cell.text = ""
            cp = cell.paragraphs[0]
            cp.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
            run = cp.add_run(txt)
            set_run_font(run, size=10)
    doc.add_paragraph()


def extract_postman_requests(collection_path: Path) -> list[dict[str, str]]:
    """Recorre la colección v2.1 y devuelve filas para documentación de pruebas."""
    if not collection_path.is_file():
        return []
    data = json.loads(collection_path.read_text(encoding="utf-8"))
    out: list[dict[str, str]] = []

    def walk(items: list, folder: str) -> None:
        for it in items:
            if "item" in it:
                walk(it["item"], it.get("name", ""))
            elif "request" in it:
                req = it["request"]
                method = str(req.get("method", "?"))
                url = req.get("url")
                if isinstance(url, dict):
                    parts = url.get("path") or []
                    path_str = "/" + "/".join(str(p).strip("{}") for p in parts) if parts else url.get("raw", "")
                else:
                    path_str = str(url)
                path_str = path_str.replace("{{base_url}}", "").strip()
                if path_str.startswith("//"):
                    path_str = path_str[1:]
                out.append(
                    {
                        "folder": folder,
                        "name": it.get("name", "?"),
                        "method": method,
                        "path": path_str or "—",
                    }
                )

    walk(data.get("item", []), "")
    return out


def criterio_aceptacion(method: str, path: str, name: str) -> str:
    p, m, n = path.lower(), method.upper(), name.lower()
    if "auth/login" in p or n == "login":
        return "HTTP 200; success true; accessToken y refreshToken en data."
    if "auth/refresh" in p:
        return "HTTP 200; success true; nuevo accessToken."
    if "auth/logout" in p:
        return "HTTP 200 o 204; sesión invalidada según implementación."
    if m == "GET":
        return "HTTP 200; success true; data con recurso o lista/página."
    if m == "POST":
        return "HTTP 200 o 201; success true; recurso creado o confirmación."
    if m == "PUT":
        return "HTTP 200; success true; recurso actualizado."
    if m in ("DELETE", "PATCH"):
        return "HTTP 200; success true; operación confirmada."
    return "HTTP 2xx; cuerpo ApiResponse coherente con la operación."


def add_anexo_placeholder(doc: Document, titulo: str) -> None:
    add_heading_h3(doc, apa_title_case(titulo))
    add_normal_paragraph(
        doc,
        "[Espacio reservado para el contenido del anexo. Incluir aquí tablas, figuras o instrumentos "
        "según lo acordado con el asesor. Los anexos se numeran y titulan para facilitar su localización; "
        "verifique la normativa institucional sobre numeración de páginas en anexos.]",
    )


def build_document() -> Document:
    doc = Document()
    for section in doc.sections:
        configure_section(section)

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT

    add_heading_centered(doc, "Capítulo V. Pruebas De Calidad De Software")

    add_heading_h2(doc, "Pruebas Unitarias")
    add_normal_paragraph(
        doc,
        "Las pruebas unitarias son un conjunto de pruebas automatizadas que se realizan para verificar el "
        "funcionamiento correcto de cada unidad individual de código, como funciones o métodos. Estas pruebas "
        "se ejecutan de forma aislada y garantizan que cada unidad funcione según lo esperado, lo que ayuda a "
        "prevenir errores y mejorar la calidad del software. En el marco del sistema para la licorería Chilalo "
        "Shot, la prioridad documentada en el proyecto ha recaído en pruebas manuales e integrales sobre la API "
        "y los clientes; no obstante, las **pruebas unitarias** constituyen la base recomendable para asegurar "
        "servicios críticos (por ejemplo, cálculo de totales, descuentos, stock y reglas de fidelización) en el "
        "backend Spring Boot. Su realización objetiva puede resumirse indicando qué clases o métodos se cubren, "
        "con qué framework se ejecutan (p. ej. JUnit 5 con Spring Boot Test) y un muestreo de resultados "
        "(pasó / falló). El detalle de casos, datos de prueba y capturas puede concentrarse en un **anexo** "
        "para no sobrecargar el cuerpo del informe.",
    )

    add_heading_h2(doc, "Pruebas Integrales")
    add_normal_paragraph(
        doc,
        "Las pruebas integrales de software verifican la correcta interacción entre diferentes unidades de código "
        "ya integradas. Se enfocan en cómo los distintos módulos del software funcionan juntos como un sistema "
        "completo, simulando escenarios reales de uso. En el proyecto de la licorería Chilalo Shot, el plan de "
        "pruebas integrales sobre el backend se deriva de la colección Postman "
        "«LICORERIA_BACKEND.postman_collection.json», ubicada en la raíz del repositorio del sistema. Dicha "
        "colección agrupa las peticiones HTTP por módulos (autenticación, productos, clientes, ventas, compras, "
        "categorías, proveedores, inventario, promociones, packs, facturación y usuarios) y sirve como catálogo "
        "reproducible de casos de prueba contra la API REST. Complementariamente puede contrastarse con la "
        "documentación Swagger y con pruebas manuales sobre el frontend React y la app Android.",
    )

    add_heading_h3(doc, "Metodología Y Herramienta")
    add_normal_paragraph(
        doc,
        "Se importa la colección en Postman (o cliente compatible) y se configura un entorno con la variable "
        "base_url (por ejemplo http://localhost:8080). La petición «Login» incluye un script de prueba que, ante "
        "respuesta exitosa, almacena access_token y refresh_token en el entorno; el resto de solicitudes usa el "
        "encabezado Authorization: Bearer {{access_token}}. La ejecución puede realizarse carpeta por carpeta "
        "(Collection Runner) o de forma manual siguiendo el orden lógico: primero autenticación, luego catálogos "
        "y finalmente transacciones (ventas, compras, facturación). Los resultados deben registrarse de forma "
        "objetiva (código HTTP, cuerpo JSON, fecha); las capturas detalladas pueden ubicarse en el Anexo 1.",
    )

    add_heading_h3(doc, "Entorno De Prueba (Referencia)")
    add_bullet_list(
        doc,
        [
            "Backend: Spring Boot en ejecución (puerto 8080 por defecto) con base PostgreSQL y migraciones Flyway aplicadas.",
            "Credenciales de ejemplo en la colección: usuario admin / contraseña Admin123! (ajustar si el entorno difiere).",
            "Herramientas: Postman Desktop, Newman para ejecución por línea de comandos (opcional) o repetición con cURL.",
        ],
    )

    postman_rows = extract_postman_requests(POSTMAN_COLLECTION)
    add_heading_h3(doc, "Catálogo De Casos De Prueba (Derivado De La Colección Postman)")
    if postman_rows:
        add_normal_paragraph(
            doc,
            "La tabla 1 lista cada petición de la colección como un caso de prueba integral PI-xx. La columna "
            "«Resultado observado» debe completarse al ejecutar las pruebas (o consignarse «Cumple» / «No cumple» "
            "más hallazgo en anexo).",
        )
        table_data: list[list[str]] = []
        for i, r in enumerate(postman_rows, start=1):
            crit = criterio_aceptacion(r["method"], r["path"], r["name"])
            table_data.append(
                [
                    f"PI-{i:02d}",
                    r["folder"],
                    r["name"],
                    r["method"],
                    r["path"],
                    crit,
                    "—",
                ]
            )
        add_table(
            doc,
            "Tabla 1. Catálogo De Casos De Prueba Integral Basado En Licorería Backend Api (Postman)",
            [
                "ID",
                "Módulo",
                "Caso (nombre en Postman)",
                "Método",
                "Endpoint",
                "Criterio de aceptación esperado",
                "Resultado observado",
            ],
            table_data,
        )
    else:
        add_normal_paragraph(
            doc,
            "No se pudo leer el archivo LICORERIA_BACKEND.postman_collection.json en la ruta esperada. "
            "Copie la colección junto a generar_cap_v_taap.py y vuelva a generar el documento.",
        )

    add_heading_h2(doc, "Conclusiones Y Recomendaciones")

    add_heading_h3(doc, "Conclusión 1")
    add_normal_paragraph(
        doc,
        "El desarrollo de una solución integral (backend, panel web y aplicación móvil) para la gestión de la "
        "licorería Chilalo Shot demuestra que una arquitectura por capas con **API REST** y autenticación "
        "**JWT** es adecuada para separar responsabilidades, facilitar el mantenimiento y permitir que varios "
        "clientes consuman la misma lógica de negocio.",
    )

    add_heading_h3(doc, "Recomendación 1")
    add_normal_paragraph(
        doc,
        "Se recomienda formalizar un plan de pruebas que combine pruebas unitarias en servicios críticos del "
        "backend con pruebas integrales reproducibles: mantener actualizada la colección "
        "LICORERIA_BACKEND.postman_collection.json, ejecutar regresión periódica (Postman Runner o Newman) "
        "antes de cada entrega y archivar resultados en el repositorio o en anexos del trabajo.",
    )

    add_heading_h3(doc, "Conclusión 2")
    add_normal_paragraph(
        doc,
        "La persistencia en **PostgreSQL** con migraciones **Flyway** y el uso de un esquema relacional alineado "
        "al dominio (ventas, inventario, clientes, compras) contribuye a la trazabilidad de las operaciones y "
        "reduce el riesgo de inconsistencias entre ambientes de desarrollo y despliegue.",
    )

    add_heading_h3(doc, "Recomendación 2")
    add_normal_paragraph(
        doc,
        "Se recomienda mantener documentación viva de la API (Swagger), respaldos periódicos de la base de datos "
        "y credenciales de integración (p. ej. SUNAT en ambiente de pruebas) fuera del código fuente, siguiendo "
        "buenas prácticas de seguridad.",
    )

    add_heading_h3(doc, "Conclusión 3")
    add_normal_paragraph(
        doc,
        "La experiencia de usuario en el **POS** y en los módulos de inventario y reportes condiciona la adopción "
        "del sistema en el día a día; la coherencia entre reglas en servidor y validaciones en cliente es "
        "determinante para evitar frustraciones y errores operativos.",
    )

    add_heading_h3(doc, "Recomendación 3")
    add_normal_paragraph(
        doc,
        "Se recomienda capacitar al personal en el uso del sistema, establecer roles claros (**ADMIN** vs "
        "**VENDEDOR**) y programar revisiones periódicas de stock, precios y promociones para que la herramienta "
        "siga alineada con la realidad del negocio.",
    )

    add_heading_h2(doc, "Referencias Bibliográficas")
    refs = [
        "Flores, E., Valenzuela, K. y Proleón, Ch. (2020). Guía de citas y referencias basado en la norma de estilo APA (7.ª ed. en inglés): Aprobada por Consejo Universitario el 8 de julio de 2020. Universidad del Pacífico. https://up-pe.libguides.com/ld.php?content_id=55976859",
        "Pontificia Universidad Católica del Perú [PUCP]. (2020, 1 de diciembre). Versión resumida de Normas APA 7 edición. Blog de la Maestría y Doctorado PUCP. http://blog.pucp.edu.pe/blog/maestriaeducacion/2020/12/01/version-resumida-apa-7-edicion/",
        "Sánchez, C. (2019, 8 de febrero). Normas APA – 7ma (séptima) edición. Normas APA actualizadas (7.ª edición). https://normas-apa.org/",
        "Universidad San Martín de Porres. (2020). Manual para la elaboración de tesis. USMP. https://www.usmp.edu.pe/iced/pdfs/manual-apa.pdf",
        "Vicerrectoría de investigación. (2021). Guía de estilo Editorial. Universidad San Ignacio de Loyola (USIL). https://378236d7-2c94-4012-bcef-3162534dcaa7.filesusr.com/ugd/ad1a884be54e4af5b34d9ca6a72a8022b6cbe7.pdf",
    ]
    for r in refs:
        add_ref_entry(doc, r)

    doc.add_page_break()

    add_heading_h2(doc, "Anexos")
    add_normal_paragraph(
        doc,
        "Los anexos son documentos que se incluyen como apoyo o complemento del trabajo final. Usualmente "
        "comprenden listas, tablas y figuras muy extensas. Estos anexos deben estar ordenados y numerados, e "
        "incluir el título de cada anexo para su pronta verificación. Según la normativa institucional, los "
        "anexos pueden no llevar numeración de páginas en el cuerpo principal del documento ni incluirse en el "
        "conteo de páginas del texto central; confirme esta regla con su asesor o reglamento de titulación.",
    )

    add_heading_h3(doc, "Anexo 1: Matriz De Servicios Y Evidencias De Pruebas Api")
    add_normal_paragraph(
        doc,
        "[Incluir la matriz de servicios si el programa la exige. Para pruebas integrales, anexar capturas de "
        "Postman por módulo (request/response), exportación de resultados del Collection Runner o salida de "
        "Newman, fecha y versión del backend probada. Puede vincularse cada captura al identificador PI-xx de la "
        "tabla 1.]",
    )
    add_anexo_placeholder(doc, "Anexo 2: Instrumento De Recolección De Datos")
    add_anexo_placeholder(doc, "Anexo 3: Ficha De Validación De Instrumento")
    add_anexo_placeholder(doc, "Anexo 4: Validación De Expertos")

    return doc


def main() -> None:
    doc = build_document()
    doc.save(OUT_FILE)
    print(f"Guardado: {OUT_FILE}")


if __name__ == "__main__":
    main()
