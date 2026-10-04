from docx import Document
from docx.shared import Pt


def add_title(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(16)


def add_h1(doc: Document, text: str) -> None:
    doc.add_heading(text, level=1)


def add_h2(doc: Document, text: str) -> None:
    doc.add_heading(text, level=2)


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def add_table(doc: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = h
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = value


def main() -> None:
    doc = Document()

    add_title(doc, "RETROSPECTIVA DEL SPRINT")
    doc.add_paragraph("Sprint 3 - Sistema de Gestion para Licoreria Chilalo Shot")
    doc.add_paragraph("Proyecto: Sistema Web y App Movil - Chilalo Shot")
    doc.add_paragraph("Sprint: Sprint 3")
    doc.add_paragraph("Fecha: 07 de abril de 2026")
    doc.add_paragraph("Facilitador: Equipo de desarrollo")
    doc.add_paragraph("Participantes: Desarrollador, propietario del negocio y docente evaluador")
    doc.add_paragraph("Duracion: 1 hora")
    doc.add_paragraph("Documento de uso interno y academico")

    add_h1(doc, "Tabla de Contenidos")
    toc_items = [
        "Tabla de Contenidos",
        "1. Introduccion: Que es una Retrospectiva del Sprint?",
        "1.1 Principios Fundamentales",
        "1.2 Formato de esta Retrospectiva",
        "2. Contexto del Sprint 3",
        "2.1 Metricas Clave del Sprint 3",
        "2.2 Estado de las Historias de Usuario",
        "2.3 Interpretacion de los Datos",
        "3. Que Salio Bien? (Glad)",
        "4. Que No Salio Bien? (Sad / Mad)",
        "4.1 Problemas Identificados",
        "4.2 Diagrama de Causa y Efecto",
        "5. Analisis Starfish (Estrella de Mar)",
        "5.1 Seguir Haciendo (Keep Doing)",
        "5.2 Hacer Mas (More Of)",
        "5.3 Empezar a Hacer (Start Doing)",
        "5.4 Hacer Menos (Less Of)",
        "5.5 Dejar de Hacer (Stop Doing)",
        "6. Plan de Accion Post-Entrega",
        "6.1 Acciones de Mejora Comprometidas",
        "6.2 Trabajo de Continuidad: Siguientes Iteraciones",
        "7. Analisis de Velocity y Cierre del Proyecto",
        "7.1 Velocity del Sprint 3",
        "7.2 Proyeccion de Evolucion",
        "8. Revision de la Definicion de Terminado (DoD)",
        "8.1 DoD Sprint 3 vs DoD Operativo",
        "9. Indicadores de Salud del Equipo",
        "10. Conclusiones y Compromisos",
        "10.1 Resumen Ejecutivo",
        "10.2 Los 3 Compromisos Principales del Equipo",
        "10.3 Seguimiento",
    ]
    for item in toc_items:
        doc.add_paragraph(item)

    add_h1(doc, "1. Introduccion: Que es una Retrospectiva del Sprint?")
    doc.add_paragraph(
        "La retrospectiva final revisa como se trabajo durante el Sprint 3, identifica aprendizajes "
        "clave del cierre del proyecto y define compromisos de continuidad para la operacion."
    )
    add_h2(doc, "1.1 Principios Fundamentales")
    add_bullets(
        doc,
        [
            "Transparencia: revisar logros y dificultades con evidencia.",
            "Inspeccion: usar datos del sprint para evaluar resultados.",
            "Adaptacion: proponer mejoras para evolucion y mantenimiento.",
        ],
    )
    add_h2(doc, "1.2 Formato de esta Retrospectiva")
    add_bullets(
        doc,
        [
            "Revision de metricas del Sprint 3.",
            "Analisis de aciertos y dificultades.",
            "Starfish para consolidar practicas del equipo.",
            "Plan de continuidad post-entrega.",
        ],
    )

    add_h1(doc, "2. Contexto del Sprint 3")
    doc.add_paragraph(
        "El Sprint 3 se enfoco en el valor diferencial del sistema: promociones, fidelizacion, app "
        "movil, reportes, capacidades de inteligencia de negocio, integraciones con perifericos, "
        "pruebas finales y puesta en produccion con capacitacion."
    )
    add_h2(doc, "2.1 Metricas Clave del Sprint 3")
    add_table(
        doc,
        ["Metrica", "Planificado", "Real", "Cumplimiento"],
        [
            ["Story Points comprometidos", "147 SP", "147 SP", "100%"],
            ["Story Points completados (Velocity)", "147 SP", "147 SP", "100%"],
            ["Historias de Usuario completadas", "5", "5 de 5", "100%"],
            ["Tareas completadas", "26", "26 de 26", "100%"],
            ["Horas aproximadas", "~120 h", "~120 h", "En rango"],
            ["Duracion del Sprint", "5 semanas", "5 semanas", "100%"],
        ],
    )
    add_h2(doc, "2.2 Estado de las Historias de Usuario")
    add_table(
        doc,
        ["ID", "Historia (resumen)", "Estado"],
        [
            ["HU-10", "Promociones y packs automaticos en POS", "Completada"],
            ["HU-11", "Programa de fidelizacion por puntos", "Completada"],
            ["HU-12", "App movil operativa para gestion y ventas", "Completada"],
            ["HU-13", "Reportes y dashboard para decisiones", "Completada"],
            ["HU-14", "Cierre final: IA, perifericos, despliegue y capacitacion", "Completada"],
        ],
    )
    add_h2(doc, "2.3 Interpretacion de los Datos")
    doc.add_paragraph(
        "El Sprint 3 se completo al 100% y cerro la implementacion funcional del proyecto. "
        "Se logro pasar de la operacion base a capacidades diferenciales para crecimiento del negocio."
    )

    add_h1(doc, "3. Que Salio Bien? (Glad)")
    add_table(
        doc,
        ["#", "Aspecto positivo", "Evidencia"],
        [
            ["1", "Cierre completo del backlog", "26 de 26 tareas finalizadas."],
            ["2", "Entrega de alto valor de negocio", "Promociones, fidelizacion y reportes operativos."],
            ["3", "Multiplataforma funcional", "Sistema web y app movil alineados."],
            ["4", "Cierre de implementacion real", "Despliegue y capacitacion completados."],
            ["5", "Madurez de ejecucion", "Sin historias pendientes y con enfoque en cliente."],
        ],
    )

    add_h1(doc, "4. Que No Salio Bien? (Sad / Mad)")
    add_h2(doc, "4.1 Problemas Identificados")
    add_table(
        doc,
        ["#", "Nivel", "Problema", "Causa raiz", "Impacto"],
        [
            [
                "1",
                "SAD",
                "Ajustes finales en IA y predicciones",
                "Necesidad de calibrar reglas y datos historicos",
                "Retrabajo puntual en fase de cierre",
            ],
            [
                "2",
                "SAD",
                "Pruebas en dispositivos reales Android",
                "Diferencias de rendimiento entre equipos",
                "Tiempo extra de validacion final",
            ],
            [
                "3",
                "MAD",
                "Acumulacion de tareas de cierre",
                "Convergencia de despliegue, pruebas y capacitacion",
                "Mayor presion en los ultimos dias",
            ],
        ],
    )
    add_h2(doc, "4.2 Diagrama de Causa y Efecto")
    add_bullets(
        doc,
        [
            "Causa: calibracion de modelos y reglas -> Efecto: ajustes finales en funcionalidades inteligentes.",
            "Causa: variabilidad de dispositivos reales -> Efecto: mas tiempo de prueba en app movil.",
            "Causa: varias actividades criticas al cierre -> Efecto: carga alta en la recta final.",
        ],
    )

    add_h1(doc, "5. Analisis Starfish (Estrella de Mar)")
    add_h2(doc, "5.1 Seguir Haciendo (Keep Doing)")
    add_bullets(
        doc,
        [
            "Enfocar cada entrega en valor directo para el negocio.",
            "Validar con escenarios reales antes del cierre.",
            "Mantener comunicacion continua con el cliente.",
        ],
    )
    add_h2(doc, "5.2 Hacer Mas (More Of)")
    add_bullets(
        doc,
        [
            "Pruebas tempranas en equipos reales de uso final.",
            "Documentacion funcional para usuarios no tecnicos.",
            "Mini demos por modulo durante el sprint.",
        ],
    )
    add_h2(doc, "5.3 Empezar a Hacer (Start Doing)")
    add_bullets(
        doc,
        [
            "Plan de mantenimiento trimestral post-entrega.",
            "Tablero de mejora continua con feedback del cliente.",
            "Monitoreo basico de uso para decisiones de evolucion.",
        ],
    )
    add_h2(doc, "5.4 Hacer Menos (Less Of)")
    add_bullets(
        doc,
        [
            "Concentrar actividades criticas en la ultima semana.",
            "Ajustes de ultimo minuto sin ventana de estabilizacion.",
        ],
    )
    add_h2(doc, "5.5 Dejar de Hacer (Stop Doing)")
    add_bullets(
        doc,
        [
            "Posponer validaciones en dispositivos reales para el final.",
            "Cerrar sprint sin checklist de operacion post-entrega.",
        ],
    )

    add_h1(doc, "6. Plan de Accion Post-Entrega")
    add_h2(doc, "6.1 Acciones de Mejora Comprometidas")
    add_table(
        doc,
        ["#", "Accion", "Responsable", "Prioridad", "Fecha limite", "Metrica de exito"],
        [
            [
                "1",
                "Definir plan de soporte y mantenimiento",
                "Equipo",
                "Alta",
                "30 dias post-entrega",
                "Incidencias atendidas con tiempos definidos",
            ],
            [
                "2",
                "Consolidar manual funcional para administrador y vendedor",
                "Desarrollador",
                "Alta",
                "2 semanas post-entrega",
                "Manual validado por usuario final",
            ],
            [
                "3",
                "Programar ciclo de mejoras con feedback real",
                "Equipo + Cliente",
                "Media",
                "Mensual",
                "Backlog de mejoras priorizado y actualizado",
            ],
        ],
    )
    add_h2(doc, "6.2 Trabajo de Continuidad: Siguientes Iteraciones")
    add_table(
        doc,
        ["Linea de continuidad", "Estado", "Enfoque siguiente"],
        [
            ["Proyecto base implementado", "Cerrado", "Entrar a etapa de mejora continua"],
            ["Optimizaciones IA y reportes", "Abierto", "Ajustes por comportamiento real del negocio"],
            ["Escalamiento movil", "Abierto", "Mejoras de experiencia y rendimiento"],
        ],
    )

    add_h1(doc, "7. Analisis de Velocity y Cierre del Proyecto")
    add_h2(doc, "7.1 Velocity del Sprint 3")
    add_table(
        doc,
        ["Concepto", "Valor"],
        [
            ["Velocity planificada", "147 SP"],
            ["Velocity real", "147 SP"],
            ["Ratio de cumplimiento", "100%"],
            ["Lectura", "Cierre de sprint con alto control y foco en entrega"],
        ],
    )
    add_h2(doc, "7.2 Proyeccion de Evolucion")
    doc.add_paragraph(
        "Con el proyecto funcionalmente cerrado, la siguiente etapa se centra en estabilidad, "
        "soporte y mejoras incrementales basadas en uso real del negocio."
    )

    add_h1(doc, "8. Revision de la Definicion de Terminado (DoD)")
    add_h2(doc, "8.1 DoD Sprint 3 vs DoD Operativo")
    add_table(
        doc,
        ["DoD Sprint 3", "DoD Operativo Propuesto"],
        [
            [
                "Funcionalidades implementadas y demostradas",
                "Funcionalidades + adopcion por usuarios en operacion diaria",
            ],
            [
                "Pruebas de cierre por modulo",
                "Pruebas periodicas de regresion y control de incidentes",
            ],
            [
                "Capacitacion y entrega final",
                "Capacitacion continua y actualizacion de manuales",
            ],
        ],
    )

    add_h1(doc, "9. Indicadores de Salud del Equipo")
    add_table(
        doc,
        ["Indicador", "Puntuacion (1-5)", "Comentario"],
        [
            ["Claridad del objetivo", "5.0", "Objetivo final del sprint claramente definido."],
            ["Colaboracion", "4.9", "Alto nivel de coordinacion en cierre."],
            ["Ritmo de trabajo", "4.5", "Sostenido, con mayor carga al final."],
            ["Confianza en la entrega", "5.0", "Cumplimiento completo del backlog."],
            ["Satisfaccion general", "4.9", "Proyecto cerrado con resultado positivo."],
        ],
    )

    add_h1(doc, "10. Conclusiones y Compromisos")
    add_h2(doc, "10.1 Resumen Ejecutivo")
    doc.add_paragraph(
        "El Sprint 3 cerro exitosamente el proyecto con el 100% del backlog completado. "
        "Se entregaron capacidades diferenciales de negocio, se realizo la puesta en produccion "
        "y se dejo al cliente capacitado para la operacion."
    )
    add_h2(doc, "10.2 Los 3 Compromisos Principales del Equipo")
    add_bullets(
        doc,
        [
            "Mantener soporte post-entrega con tiempos de respuesta definidos.",
            "Priorizar mejoras por impacto real en la operacion del negocio.",
            "Sostener lenguaje claro y orientado al usuario final en toda comunicacion.",
        ],
    )
    add_h2(doc, "10.3 Seguimiento")
    doc.add_paragraph(
        "El seguimiento se realizara con revisiones periodicas de uso, incidencias y mejoras, "
        "asegurando la evolucion del sistema despues del cierre del proyecto."
    )

    out = (
        r"c:\Users\Usuario\Desktop\Liquor-Store-Piura-master\Liquor-Store-Piura-master"
        r"\Retrospectiva_Sprint3_CHILALO_SHOT.docx"
    )
    doc.save(out)
    print(out)


if __name__ == "__main__":
    main()
