from docx import Document
from docx.shared import Pt


def add_title(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(16)


def add_h1(doc: Document, text: str) -> None:
    doc.add_heading(text, level=1)


def add_h2(doc: Document, text: str) -> None:
    doc.add_heading(text, level=2)


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def add_table(doc: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = h
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = value


def main() -> None:
    doc = Document()

    add_title(doc, "RETROSPECTIVA DEL SPRINT")
    doc.add_paragraph("Sprint 1 - Sistema de Gestion para Licoreria Chilalo Shot")
    doc.add_paragraph("Proyecto: Sistema Web y App Movil - Chilalo Shot")
    doc.add_paragraph("Sprint: Sprint 1")
    doc.add_paragraph("Fecha: 07 de abril de 2026")
    doc.add_paragraph("Facilitador: Equipo de desarrollo")
    doc.add_paragraph("Participantes: Desarrollador, propietario del negocio y docente evaluador")
    doc.add_paragraph("Duracion: 1 hora")
    doc.add_paragraph("Documento de uso interno y academico")

    add_h1(doc, "Tabla de Contenidos")
    toc_items = [
        "Tabla de Contenidos",
        "1. Introduccion: Que es una Retrospectiva del Sprint?",
        "1.1 Principios Fundamentales",
        "1.2 Formato de esta Retrospectiva",
        "2. Contexto del Sprint 1",
        "2.1 Metricas Clave del Sprint 1",
        "2.2 Estado de las Historias de Usuario",
        "2.3 Interpretacion de los Datos",
        "3. Que Salio Bien? (Glad)",
        "4. Que No Salio Bien? (Sad / Mad)",
        "4.1 Problemas Identificados",
        "4.2 Diagrama de Causa y Efecto",
        "5. Analisis Starfish (Estrella de Mar)",
        "5.1 Seguir Haciendo (Keep Doing)",
        "5.2 Hacer Mas (More Of)",
        "5.3 Empezar a Hacer (Start Doing)",
        "5.4 Hacer Menos (Less Of)",
        "5.5 Dejar de Hacer (Stop Doing)",
        "6. Plan de Accion para el Sprint 2",
        "6.1 Acciones de Mejora Comprometidas",
        "6.2 Historias No Completadas: Estrategia para Sprint 2",
        "7. Analisis de Velocity y Proyeccion",
        "7.1 Velocity del Sprint 1",
        "7.2 Proyeccion del Proyecto",
        "8. Revision de la Definicion de Terminado (DoD)",
        "8.1 DoD Sprint 1 vs DoD Propuesto Sprint 2",
        "9. Indicadores de Salud del Equipo",
        "10. Conclusiones y Compromisos",
        "10.1 Resumen Ejecutivo",
        "10.2 Los 3 Compromisos Principales del Equipo",
        "10.3 Seguimiento",
    ]
    for item in toc_items:
        doc.add_paragraph(item)

    add_h1(doc, "1. Introduccion: Que es una Retrospectiva del Sprint?")
    doc.add_paragraph(
        "Esta retrospectiva resume que funciono, que dificultades aparecieron y que acciones "
        "tomaremos para mejorar en el siguiente sprint. El enfoque es practico: mejorar el trabajo "
        "del equipo y asegurar resultados claros para el negocio."
    )
    add_h2(doc, "1.1 Principios Fundamentales")
    add_bullets(
        doc,
        [
            "Transparencia: todos pueden opinar con confianza y respeto.",
            "Inspeccion: se revisan datos reales, no solo percepciones.",
            "Adaptacion: se acuerdan mejoras concretas para el siguiente sprint.",
        ],
    )
    add_h2(doc, "1.2 Formato de esta Retrospectiva")
    add_bullets(
        doc,
        [
            "Revision de resultados reales del Sprint 1.",
            "Analisis de lo positivo y lo mejorable.",
            "Definicion de acciones concretas para Sprint 2.",
            "Lenguaje orientado al cliente y al valor del negocio.",
        ],
    )

    add_h1(doc, "2. Contexto del Sprint 1")
    doc.add_paragraph(
        "El Sprint 1 tuvo como meta dejar lista la base del sistema: acceso seguro, recuperacion de "
        "contrasena, gestion de usuarios, categorias y productos, y configuracion general del negocio."
    )
    add_h2(doc, "2.1 Metricas Clave del Sprint 1")
    add_table(
        doc,
        ["Metrica", "Planificado", "Real", "Cumplimiento"],
        [
            ["Tareas del backlog", "13", "13", "100%"],
            ["Tareas adicionales", "0", "2", "Superado"],
            ["Horas aproximadas", "~80 h", "~80 h", "En rango"],
            ["Duracion del sprint", "4 semanas", "4 semanas", "100%"],
        ],
    )
    add_h2(doc, "2.2 Estado de las Historias de Usuario")
    add_table(
        doc,
        ["ID", "Historia (resumen)", "Estado"],
        [
            ["HU-01", "Base ordenada del sistema y requerimientos iniciales", "Completada"],
            ["HU-02", "Ingreso seguro, recuperacion de contrasena y control de accesos", "Completada"],
            ["HU-03", "Gestion de categorias y productos con busqueda/filtros", "Completada"],
            ["HU-04", "Configuracion de datos del negocio", "Completada"],
            ["HU-05", "Adelantos de la app movil", "Completada (adicional)"],
        ],
    )
    add_h2(doc, "2.3 Interpretacion de los Datos")
    doc.add_paragraph(
        "El sprint se cumplio al 100% en lo planificado y ademas se entregaron 2 tareas extra. "
        "Se logro una base estable para avanzar al Sprint 2 con foco en operacion diaria (ventas e inventario)."
    )

    add_h1(doc, "3. Que Salio Bien? (Glad)")
    add_table(
        doc,
        ["#", "Aspecto positivo", "Evidencia"],
        [
            ["1", "Objetivo del sprint claro", "Se completo la base funcional esperada."],
            ["2", "Avance completo del backlog", "13 de 13 tareas terminadas."],
            ["3", "Valor visible para el cliente", "Login, usuarios, categorias y productos ya operativos."],
            ["4", "Buena comunicacion con cliente", "Validaciones y feedback recibidos durante el sprint."],
            ["5", "Trabajo adicional entregado", "Se adelantaron 2 tareas moviles no planificadas."],
        ],
    )

    add_h1(doc, "4. Que No Salio Bien? (Sad / Mad)")
    add_h2(doc, "4.1 Problemas Identificados")
    add_table(
        doc,
        ["#", "Nivel", "Problema", "Causa raiz", "Impacto"],
        [
            [
                "1",
                "SAD",
                "Incidente de conexion entre web y servidor al inicio",
                "Ajuste inicial de permisos de comunicacion entre aplicaciones",
                "Retraso aproximado de 3 horas en la semana 1",
            ],
            [
                "2",
                "SAD",
                "Algunas tareas crecieron en alcance",
                "Se agregaron mejoras utiles durante el desarrollo",
                "Mayor esfuerzo de implementacion en modulo de productos",
            ],
            [
                "3",
                "MAD",
                "Riesgo de lenguaje demasiado tecnico para cliente",
                "Documentos y descripciones con enfoque interno",
                "Puede dificultar comprension en revisiones",
            ],
        ],
    )

    add_h2(doc, "4.2 Diagrama de Causa y Efecto")
    add_bullets(
        doc,
        [
            "Causa: ajustes tecnicos iniciales de entorno -> Efecto: retraso corto en semana 1.",
            "Causa: mejoras agregadas durante desarrollo -> Efecto: mas carga en tareas centrales.",
            "Causa: redaccion tecnica -> Efecto: menor claridad para usuarios no tecnicos.",
        ],
    )

    add_h1(doc, "5. Analisis Starfish (Estrella de Mar)")
    add_h2(doc, "5.1 Seguir haciendo (Keep Doing)")
    add_bullets(
        doc,
        [
            "Validar avances frecuentes con el cliente.",
            "Mantener foco en funcionalidades utiles para el negocio.",
            "Cerrar tareas con demostracion funcional.",
        ],
    )
    add_h2(doc, "5.2 Hacer mas (More Of)")
    add_bullets(
        doc,
        [
            "Probar escenarios reales del negocio desde etapas tempranas.",
            "Registrar pedidos del cliente y convertirlos rapido en backlog.",
            "Revisar textos con enfoque de usuario final.",
        ],
    )
    add_h2(doc, "5.3 Empezar a hacer (Start Doing)")
    add_bullets(
        doc,
        [
            "Definir criterio 'texto entendible para cliente' antes de cada review.",
            "Agregar mini checklists de demo por historia de usuario.",
            "Planificar validaciones de calidad desde mitad de sprint.",
        ],
    )
    add_h2(doc, "5.4 Hacer menos (Less Of)")
    add_bullets(
        doc,
        [
            "Cambios de alcance no priorizados dentro del sprint.",
            "Descripciones extensas con terminos tecnicos en presentaciones.",
        ],
    )
    add_h2(doc, "5.5 Dejar de hacer (Stop Doing)")
    add_bullets(
        doc,
        [
            "Posponer la adaptacion de lenguaje para cliente al final.",
            "Presentar siglas tecnicas sin explicacion de beneficio.",
        ],
    )

    add_h1(doc, "6. Plan de Accion para el Sprint 2")
    add_h2(doc, "6.1 Acciones de Mejora Comprometidas")
    add_table(
        doc,
        ["#", "Accion", "Responsable", "Prioridad", "Fecha limite", "Metrica de exito"],
        [
            [
                "1",
                "Usar lenguaje orientado al negocio en review y documentos",
                "Equipo",
                "Alta",
                "Inicio Sprint 2",
                "100% de historias explicadas en terminos de valor para usuario",
            ],
            [
                "2",
                "Incluir stock visible en lista de productos",
                "Desarrollador",
                "Alta",
                "Sprint 2",
                "Stock visible y validado por cliente",
            ],
            [
                "3",
                "Preparar guion de demo corto por modulo",
                "Equipo",
                "Media",
                "Sprint 2",
                "Demo fluida, sin retrabajo en reunion",
            ],
            [
                "4",
                "Validar entorno al inicio para evitar bloqueos de conexion",
                "Desarrollador",
                "Media",
                "Semana 1 Sprint 2",
                "0 bloqueos por entorno",
            ],
        ],
    )

    add_h2(doc, "6.2 Historias No Completadas: Estrategia para Sprint 2")
    add_table(
        doc,
        ["Historia", "Estado", "Estrategia para Sprint 2"],
        [
            ["Ninguna historia pendiente", "Completado", "No aplica; backlog del Sprint 1 se cerro al 100%."],
            [
                "Pedido de cliente: stock visible en lista",
                "Nueva mejora",
                "Se incorpora como prioridad alta en el backlog del Sprint 2.",
            ],
        ],
    )

    add_h1(doc, "7. Analisis de Velocity y Proyeccion")
    add_h2(doc, "7.1 Velocity del Sprint 1")
    add_table(
        doc,
        ["Concepto", "Valor"],
        [
            ["Cumplimiento del backlog", "100% (13/13 tareas)"],
            ["Tareas adicionales", "2"],
            ["Horas aproximadas", "~80 h"],
            ["Lectura de capacidad", "Capacidad adecuada para lo planificado en Sprint 1"],
        ],
    )
    add_h2(doc, "7.2 Proyeccion del Proyecto")
    doc.add_paragraph(
        "Con la base del sistema ya operativa en Sprint 1, Sprint 2 se enfoca en la operacion diaria "
        "de la licoreria: ventas, inventario, compras, caja y comprobantes. Se espera mantener alto "
        "cumplimiento priorizando valor directo para el negocio."
    )

    add_h1(doc, "8. Revision de la Definicion de Terminado (DoD)")
    add_h2(doc, "8.1 DoD Sprint 1 vs DoD Propuesto Sprint 2")
    add_table(
        doc,
        ["DoD actual", "DoD propuesto para Sprint 2"],
        [
            [
                "Funciona en desarrollo y se demuestra en review",
                "Funciona + validacion de flujo real de negocio + explicacion simple para cliente",
            ],
            [
                "Pruebas basicas de funcionamiento",
                "Pruebas funcionales con casos reales de uso del administrador/vendedor",
            ],
            [
                "Cierre tecnico de tarea",
                "Cierre tecnico + evidencia funcional (captura o demo guiada)",
            ],
        ],
    )

    add_h1(doc, "9. Indicadores de Salud del Equipo")
    add_table(
        doc,
        ["Indicador", "Puntuacion (1-5)", "Comentario"],
        [
            ["Claridad del objetivo", "4.8", "Objetivo de sprint bien definido y entendible."],
            ["Colaboracion", "4.7", "Buena coordinacion y respuesta a bloqueos."],
            ["Ritmo de trabajo", "4.3", "Adecuado, con pico por incidente tecnico inicial."],
            ["Confianza en la entrega", "4.8", "Backlog completado al 100%."],
            ["Satisfaccion general", "4.7", "Resultado positivo y base solida para Sprint 2."],
        ],
    )

    add_h1(doc, "10. Conclusiones y Compromisos")
    add_h2(doc, "10.1 Resumen Ejecutivo")
    doc.add_paragraph(
        "El Sprint 1 fue exitoso: se completo todo lo planificado y se entregaron dos avances "
        "adicionales. El sistema ya permite operaciones clave de administracion y deja preparada la "
        "continuidad hacia el modulo operativo del Sprint 2."
    )
    add_h2(doc, "10.2 Los 3 Compromisos Principales del Equipo")
    add_bullets(
        doc,
        [
            "Mantener cada entrega explicada en terminos simples para el cliente.",
            "Priorizar funcionalidades con impacto directo en la operacion diaria.",
            "Reducir bloqueos iniciales con checklist de entorno antes de desarrollar.",
        ],
    )
    add_h2(doc, "10.3 Seguimiento")
    doc.add_paragraph(
        "Estos compromisos se revisaran en la retrospectiva del Sprint 2 y se mediran con evidencia "
        "de cumplimiento durante las demos y validaciones con el cliente."
    )

    out = (
        r"c:\Users\Usuario\Desktop\Liquor-Store-Piura-master\Liquor-Store-Piura-master"
        r"\Retrospectiva_Sprint1_CHILALO_SHOT_v2.docx"
    )
    doc.save(out)
    print(out)


if __name__ == "__main__":
    main()
