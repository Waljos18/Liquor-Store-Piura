from docx import Document
from docx.shared import Pt, RGBColor, Cm, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import copy

doc = Document()

# ── Márgenes ──────────────────────────────────────────────────────────────────
for section in doc.sections:
    section.top_margin    = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin   = Cm(3)
    section.right_margin  = Cm(2.5)

# ── Colores corporativos ───────────────────────────────────────────────────────
AZUL       = RGBColor(0x1E, 0x40, 0x8A)   # azul oscuro títulos
AZUL_CLARO = RGBColor(0x2E, 0x74, 0xB5)   # subtítulos
VERDE      = RGBColor(0x1F, 0x7A, 0x4A)
ROJO       = RGBColor(0xC0, 0x39, 0x2B)
NARANJA    = RGBColor(0xE6, 0x7E, 0x22)
GRIS_CELL  = RGBColor(0xBD, 0xD7, 0xEE)   # cabecera tabla azul claro
BLANCO     = RGBColor(0xFF, 0xFF, 0xFF)
GRIS_FILA  = RGBColor(0xF2, 0xF2, 0xF2)

# ── Helpers ────────────────────────────────────────────────────────────────────
def set_cell_bg(cell, hex_color: str):
    """Pinta el fondo de una celda."""
    tc   = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd  = OxmlElement('w:shd')
    shd.set(qn('w:val'),   'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'),  hex_color)
    tcPr.append(shd)

def set_cell_border(cell, **kwargs):
    """Borde individual por lado: top, bottom, left, right."""
    tc   = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    for side, color in kwargs.items():
        border = OxmlElement(f'w:{side}')
        border.set(qn('w:val'),   'single')
        border.set(qn('w:sz'),    '6')
        border.set(qn('w:space'), '0')
        border.set(qn('w:color'), color)
        tcBorders.append(border)
    tcPr.append(tcBorders)

def heading1(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after  = Pt(6)
    run = p.add_run(text)
    run.bold      = True
    run.font.size = Pt(16)
    run.font.color.rgb = AZUL
    # línea divisoria debajo
    pPr  = p._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bot  = OxmlElement('w:bottom')
    bot.set(qn('w:val'),   'single')
    bot.set(qn('w:sz'),    '6')
    bot.set(qn('w:space'), '1')
    bot.set(qn('w:color'), '1E408A')
    pBdr.append(bot)
    pPr.append(pBdr)

def heading2(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after  = Pt(4)
    run = p.add_run(text)
    run.bold      = True
    run.font.size = Pt(13)
    run.font.color.rgb = AZUL_CLARO

def heading3(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after  = Pt(2)
    run = p.add_run(text)
    run.bold      = True
    run.font.size = Pt(11)
    run.font.color.rgb = AZUL

def body(text, bold=False, italic=False, color=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    run.bold   = bold
    run.italic = italic
    run.font.size = Pt(10)
    if color:
        run.font.color.rgb = color

def bullet(text, level=0):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.left_indent  = Cm(0.5 + level * 0.5)
    p.paragraph_format.space_after  = Pt(2)
    run = p.add_run(text)
    run.font.size = Pt(10)

def numbered(text):
    p = doc.add_paragraph(style='List Number')
    p.paragraph_format.space_after = Pt(3)
    run = p.add_run(text)
    run.bold      = True
    run.font.size = Pt(10)

def simple_table(headers, rows, hdr_color='1E408A', alt=True):
    """Tabla genérica con encabezado coloreado."""
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    # encabezado
    for i, h in enumerate(headers):
        cell = t.rows[0].cells[i]
        set_cell_bg(cell, hdr_color)
        p    = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run  = p.add_run(h)
        run.bold = True
        run.font.size = Pt(9)
        run.font.color.rgb = BLANCO
    # filas
    for r_i, row in enumerate(rows):
        for c_i, val in enumerate(row):
            cell = t.rows[r_i + 1].cells[c_i]
            if alt and r_i % 2 == 1:
                set_cell_bg(cell, 'F2F2F2')
            p   = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if c_i > 0 else WD_ALIGN_PARAGRAPH.LEFT
            run = p.add_run(str(val))
            run.font.size = Pt(9)
    doc.add_paragraph()
    return t

def info_box(label, text, label_color='1E408A'):
    """Cuadro de principio P1/P2/P3."""
    t = doc.add_table(rows=1, cols=2)
    t.style = 'Table Grid'
    # columna etiqueta
    c0 = t.rows[0].cells[0]
    set_cell_bg(c0, label_color)
    c0.width = Cm(1.8)
    p0 = c0.paragraphs[0]
    p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r0 = p0.add_run(label)
    r0.bold = True
    r0.font.size = Pt(14)
    r0.font.color.rgb = BLANCO
    # columna texto
    c1 = t.rows[0].cells[1]
    p1 = c1.paragraphs[0]
    r1 = p1.add_run(text)
    r1.font.size = Pt(10)
    doc.add_paragraph()

# ══════════════════════════════════════════════════════════════════════════════
# PORTADA
# ══════════════════════════════════════════════════════════════════════════════
doc.add_paragraph()
doc.add_paragraph()
doc.add_paragraph()

p_title = doc.add_paragraph()
p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p_title.add_run('RETROSPECTIVA DEL SPRINT')
r.bold = True
r.font.size = Pt(26)
r.font.color.rgb = AZUL

# línea debajo del título
pPr  = p_title._p.get_or_add_pPr()
pBdr = OxmlElement('w:pBdr')
bot  = OxmlElement('w:bottom')
bot.set(qn('w:val'),   'single')
bot.set(qn('w:sz'),    '12')
bot.set(qn('w:space'), '1')
bot.set(qn('w:color'), '1E408A')
pBdr.append(bot)
pPr.append(pBdr)

doc.add_paragraph()
p_sub = doc.add_paragraph()
p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
r2 = p_sub.add_run('Sprint 2 — Sistema de Gestión Licorería Chilalo Shot')
r2.bold = True
r2.font.size = Pt(14)
r2.font.color.rgb = AZUL_CLARO

doc.add_paragraph()

# Tabla de datos de portada
cover_data = [
    ['Proyecto:',       'Sistema Web y App Móvil – Licorería Chilalo Shot'],
    ['Sprint:',         'Sprint 2'],
    ['Fecha:',          'Fin de semana 9 (Marzo 2025)'],
    ['Facilitador:',    'Alumno desarrollador (rol Scrum Master)'],
    ['Participantes:',  'Alumno desarrollador · Propietario Chilalo Shot · Docente evaluador'],
    ['Duración:',       '1 hora 30 minutos'],
]
t_cover = doc.add_table(rows=len(cover_data), cols=2)
t_cover.style = 'Table Grid'
t_cover.alignment = WD_TABLE_ALIGNMENT.CENTER
for i, (lbl, val) in enumerate(cover_data):
    c0 = t_cover.rows[i].cells[0]
    c1 = t_cover.rows[i].cells[1]
    set_cell_bg(c0, 'D9E1F2')
    p0 = c0.paragraphs[0]; r0 = p0.add_run(lbl); r0.bold=True; r0.font.size=Pt(10)
    p1 = c1.paragraphs[0]; r1 = p1.add_run(val);  r1.font.size=Pt(10)

doc.add_paragraph()
doc.add_paragraph()
p_conf = doc.add_paragraph()
p_conf.alignment = WD_ALIGN_PARAGRAPH.CENTER
rc = p_conf.add_run('Documento confidencial — Uso interno del equipo')
rc.italic = True
rc.font.size = Pt(9)
rc.font.color.rgb = RGBColor(0x70,0x70,0x70)

doc.add_page_break()

# ══════════════════════════════════════════════════════════════════════════════
# 1. INTRODUCCIÓN
# ══════════════════════════════════════════════════════════════════════════════
heading1('1. Introducción: ¿Qué es una Retrospectiva del Sprint?')
body('La Retrospectiva del Sprint es uno de los cinco eventos formales definidos en el marco de trabajo Scrum. Se realiza al finalizar cada Sprint, después del Sprint Review, y tiene como propósito fundamental que el equipo Scrum reflexione sobre su forma de trabajar e identifique mejoras concretas para el siguiente Sprint.')
body('A diferencia del Sprint Review (que se enfoca en el producto), la Retrospectiva se centra en el proceso, las personas y las herramientas. Es el momento en que el equipo se mira a sí mismo con honestidad y busca formas de ser más efectivo.')

heading2('1.1 Principios Fundamentales')
info_box('P1', 'TRANSPARENCIA\nTodos los participantes expresan con honestidad sus percepciones sobre el sprint. En un proyecto individual, esto implica autoevaluación sincera e incluir la voz del cliente (propietario).')
info_box('P2', 'INSPECCIÓN\nSe examinan los datos reales del Sprint 2 —métricas, historias, horas— y se contrasta lo planificado con lo ejecutado para identificar patrones objetivos.', label_color='1F7A4A')
info_box('P3', 'ADAPTACIÓN\nCon base en lo inspeccionado se definen mejoras concretas y medibles para el Sprint 3 en términos de proceso, herramientas y calidad del producto.', label_color='7B3F9E')

heading2('1.2 Formato de esta Retrospectiva')
body('Para esta retrospectiva del Sprint 2 del proyecto de la Licorería Chilalo Shot, se utilizó una combinación de técnicas adaptadas al contexto de un proyecto académico individual:')
bullet('Navegante-Estrellas (Starfish): para categorizar las prácticas del desarrollador en cinco dimensiones.')
bullet('Mad-Sad-Glad: para capturar el estado emocional respecto al sprint.')
bullet('Análisis cuantitativo: utilizando las métricas reales del Sprint Backlog y el Sprint Review.')
bullet('Plan de acción: con compromisos concretos y criterios de éxito medibles para el Sprint 3.')

doc.add_page_break()

# ══════════════════════════════════════════════════════════════════════════════
# 2. CONTEXTO DEL SPRINT 2
# ══════════════════════════════════════════════════════════════════════════════
heading1('2. Contexto del Sprint 2')
body('Antes de reflexionar sobre lo que salió bien o mal, es fundamental tener presentes los datos objetivos del Sprint 2. Estos números dan una base común para la discusión entre el desarrollador, el propietario del negocio y el docente evaluador.')

heading2('2.1 Métricas Clave del Sprint 2')
simple_table(
    ['Métrica', 'Planificado', 'Real', 'Cumplimiento'],
    [
        ['Story Points comprometidos',         '32 SP',      '32 SP',      '100%'],
        ['Story Points completados (Velocity)', '32 SP',      '32 SP',      '100%'],
        ['Historias de Usuario completadas',   '4',          '4 de 4',     '100%'],
        ['Tareas técnicas completadas',        '23',         '23 de 23',   '100%'],
        ['Horas estimadas',                    '100 h',      '107 h',      '+7 h (sobre lo estimado)'],
        ['Duración del Sprint',                '5 semanas',  '5 semanas',  '100%'],
    ]
)

heading2('2.2 Estado de las Historias de Usuario')
body('El siguiente cuadro muestra el estado final de cada historia de usuario al cierre del Sprint 2, incluyendo el número de tareas técnicas completadas por historia:')
simple_table(
    ['ID', 'Historia de Usuario', 'SP', 'Tareas', 'Estado', 'Observación'],
    [
        ['HU-06', 'POS web: carrito, formas de pago, vuelto, historial', '8',  '7/7', 'Completado', 'Incluye crédito/fiado y pago mixto'],
        ['HU-07', 'Facturación electrónica SUNAT (boleta, factura, RCB)', '6',  '4/4', 'Completado', 'Operativo en ambiente sandbox'],
        ['HU-08', 'Inventario: alertas, movimientos, compras, proveedores, mermas', '8',  '5/5', 'Completado', 'Stock actualiza en tiempo real'],
        ['HU-09', 'Caja, gastos, devoluciones, crédito y clientes', '10', '7/7', 'Completado', '5 módulos complementarios integrados'],
    ]
)

heading2('2.3 Interpretación de los Datos')
body('El Sprint 2 representa el mayor logro técnico del proyecto hasta ahora. Las cuatro historias de usuario se completaron al 100%, sumando los 32 story points planificados y las 23 tareas técnicas. El sistema pasó de ser una base con CRUD de productos y usuarios a un sistema completamente operativo para el negocio en su día a día.')
body('El único desvío fue en las horas: se estimaron 100 horas y se utilizaron 107, un sobrecosto del 7% que se explica principalmente por la complejidad de la integración con la API de SUNAT y por la densidad de módulos acumulados en la semana 9. Este sobrecosto no generó retraso en la entrega porque el sprint contaba con margen suficiente.')
body('El problema principal de este sprint no fue de completitud sino de densidad: agrupar 9 módulos distintos en 5 semanas para un solo desarrollador fue un desafío que se resolvió pero dejó lecciones importantes sobre la granularidad de la planificación.')

doc.add_page_break()

# ══════════════════════════════════════════════════════════════════════════════
# 3. QUÉ SALIÓ BIEN
# ══════════════════════════════════════════════════════════════════════════════
heading1('3. ¿Qué Salió Bien? (Glad)')
body('En esta sección se identificaron los aspectos positivos del Sprint 2 que deben mantenerse y reforzarse en el Sprint 3. Cada punto fue discutido en la reunión de revisión con el propietario y el docente evaluador.')
simple_table(
    ['#', 'Aspecto Positivo', 'Detalle y Evidencia'],
    [
        ['1', 'Flujo end-to-end completamente funcional',
         'El ciclo completo venta → descuento de stock → emisión de comprobante funciona sin errores. El propietario lo demostró en la Sprint Review con productos reales cargados. Valida que la arquitectura soporta la complejidad del negocio.'],
        ['2', 'Integración SUNAT operativa en sandbox',
         'Se logró generar XML en formato UBL 2.1, comprimir en ZIP y simular el envío al OSE. El PDF de boleta y factura se genera correctamente con todos los campos requeridos por SUNAT. Es el módulo técnicamente más complejo del proyecto.'],
        ['3', 'Migraciones Flyway sin errores (V8–V14)',
         'Se crearon y ejecutaron 7 migraciones nuevas sin necesidad de rollback. La estrategia de migraciones versionadas demostró ser la correcta para mantener la BD sincronizada entre entornos.'],
        ['4', 'Validación temprana con el cliente',
         'El propietario participó en dos revisiones intermedias (semana 6 y semana 8) además del Sprint Review. Esto permitió ajustar detalles del POS antes de que se acumularan.'],
        ['5', 'Uso de JPA Specification para consultas dinámicas',
         'La implementación de Specification<T> en InventarioService y VentaService permitió filtros dinámicos sin duplicar código. Esta decisión técnica facilita el módulo de reportes del Sprint 3.'],
    ],
    hdr_color='1F7A4A'
)

doc.add_page_break()

# ══════════════════════════════════════════════════════════════════════════════
# 4. QUÉ NO SALIÓ BIEN
# ══════════════════════════════════════════════════════════════════════════════
heading1('4. ¿Qué No Salió Bien? (Sad / Mad)')
body('Esta sección recoge los problemas, fricciones e impedimentos experimentados durante el Sprint 2. Cada punto se clasificó en Sad (decepcionante) o Mad (frustrante) según su impacto en el desarrollo.')

heading2('4.1 Problemas Identificados')
simple_table(
    ['#', 'Nivel', 'Problema', 'Análisis de Causa Raíz y Evidencia'],
    [
        ['1', 'MAD', 'Integración SUNAT más compleja de lo estimado',
         'El formato XML UBL 2.1 tiene reglas estrictas no documentadas del todo en la guía del OSE. Se invirtieron ~12 horas adicionales depurando el XML. Causa raíz: la estimación de HU-07 no consideró la curva de aprendizaje de la especificación.'],
        ['2', 'MAD', 'Semana 9 sobrecargada con 5 módulos simultáneos',
         'Los módulos de caja, gastos, devoluciones, crédito y clientes se agruparon en una sola semana. Aunque todos se completaron, no hubo tiempo suficiente para probar las interacciones entre módulos. Causa raíz: planificación poco granular en HU-09.'],
        ['3', 'SAD', 'Sin pruebas automatizadas en ningún módulo',
         'Todo el Sprint 2 se validó con pruebas manuales en Postman. No se escribió ningún test unitario ni de integración con JUnit. Causa raíz: el ritmo de desarrollo priorizó la entrega funcional sobre la cobertura de pruebas.'],
        ['4', 'SAD', 'El RCB (Resumen Diario de Boletas) es simulado, no real',
         'En ambiente sandbox el SunatService simula el envío con respuesta automática, pero no genera el XML del RC agrupado que SUNAT requiere en producción. Causa raíz: complejidad excedía el tiempo disponible; se acordó dejarlo para producción.'],
        ['5', 'SAD', 'Feedback del cliente llegó tarde en algunos módulos',
         'El propietario revisó el módulo de cierre de caja recién en el Sprint Review, cuando ya no había tiempo para ajustar la visualización. Causa raíz: el propietario no tenía acceso autónomo al sistema de pruebas durante el sprint.'],
    ],
    hdr_color='C0392B'
)

heading2('4.2 Diagrama de Causa y Efecto')
body('El análisis identificó que la mayoría de los problemas tienen origen en la planificación y en la falta de acceso continuo del cliente al sistema en desarrollo:')
simple_table(
    ['CAUSA RAÍZ', 'EFECTO DIRECTO', 'IMPACTO EN EL SPRINT'],
    [
        ['Subestimación de la complejidad SUNAT (UBL 2.1)', '12 horas adicionales de depuración XML', 'HU-07 consumió más tiempo; las otras HUs absorbieron el exceso'],
        ['HU-09 con 5 módulos en 1 semana', 'Ritmo insostenible en semana 9', 'Pruebas entre módulos insuficientes antes de la demo'],
        ['Desarrollo sin tests automatizados', 'Errores de regresión no detectados automáticamente', 'Mayor tiempo en revisión manual; riesgo acumulado para Sprint 3'],
        ['Cliente sin acceso autónomo al sistema de pruebas', 'Feedback tardío sobre interfaces', 'Ajustes de UI no realizados antes del Sprint Review'],
        ['RCB real no implementado', 'Módulo SUNAT incompleto para producción', 'Tarea pendiente antes de la puesta en producción'],
    ],
    hdr_color='E67E22'
)

doc.add_page_break()

# ══════════════════════════════════════════════════════════════════════════════
# 5. ANÁLISIS STARFISH
# ══════════════════════════════════════════════════════════════════════════════
heading1('5. Análisis Starfish (Estrella de Mar)')
body('La técnica Starfish organiza las reflexiones en cinco categorías que representan acciones concretas sobre las prácticas actuales. Es útil porque no solo mira lo bueno y lo malo, sino que propone acciones graduales y específicas.')

heading2('5.1 Seguir Haciendo (Keep Doing)')
body('Prácticas que funcionaron bien y deben mantenerse exactamente como están.', italic=True)
simple_table(
    ['#', 'Práctica', 'Por qué mantenerla'],
    [
        ['1', 'Revisiones intermedias con el cliente (mitad del sprint)', 'Permiten detectar errores de interpretación antes de que sea tarde. En Sprint 2 se ahorraron correcciones mayores del POS gracias a la revisión de la semana 6.'],
        ['2', 'Migraciones Flyway versionadas para todos los cambios de BD', 'Garantizan que el esquema sea reproducible. 7 migraciones aplicadas sin errores ni rollback.'],
        ['3', 'Estructura Controller → Service → Repository sin excepciones', 'Mantuvo el código organizado a pesar de la cantidad de módulos nuevos. Facilita el mantenimiento.'],
        ['4', 'Validación de endpoints con Postman antes de integrarlos al frontend', 'Detecta errores de backend antes de que el frontend los consuma, reduciendo tiempo de depuración.'],
    ],
    hdr_color='1F7A4A'
)

heading2('5.2 Hacer Más (More Of)')
body('Prácticas que ya existen pero deben intensificarse o hacerse con mayor frecuencia.', italic=True)
simple_table(
    ['#', 'Práctica', 'Cómo intensificarla'],
    [
        ['1', 'Revisiones del cliente durante el sprint', 'Pasar de 2 revisiones a 3: inicio, mitad y pre-entrega. Compartir link al sistema de pruebas para que el propietario revise desde su celular en cualquier momento.'],
        ['2', 'Documentación de decisiones técnicas complejas', 'Cuando se tome una decisión técnica importante (como la estructura XML SUNAT), registrarla en un comentario Javadoc o en DECISIONES.md para no reconstruirla después.'],
        ['3', 'Granularidad en la planificación de historias de usuario', 'Dividir historias con más de 3 módulos en sub-historias con estimación independiente para evitar semanas sobrecargadas.'],
    ],
    hdr_color='2E74B5'
)

heading2('5.3 Empezar a Hacer (Start Doing)')
body('Prácticas nuevas que no se realizaban y deben incorporarse en el Sprint 3.', italic=True)
simple_table(
    ['#', 'Nueva Práctica', 'Cómo implementarla'],
    [
        ['1', 'Pruebas automatizadas desde el inicio del sprint', 'Escribir al menos tests de integración con @SpringBootTest para los flujos críticos: crear venta, emitir boleta, calcular cierre de caja. Objetivo: cobertura mínima 50% en servicios principales.'],
        ['2', 'Pruebas de la app Android en dispositivo físico desde la primera semana', 'No dejar las pruebas del Android solo para el final. Cada pantalla nueva se prueba en el celular el mismo día que se termina.'],
        ['3', 'Checklist de criterios de aceptación antes de cerrar una tarea', 'Crear una lista de 3 a 5 criterios por historia al inicio del sprint. Una tarea no se cierra sin que todos estén verificados.'],
        ['4', 'Compartir acceso al sistema de pruebas con el propietario desde la semana 1', 'Configurar la URL del backend y frontend accesible (ngrok o similar) para feedback continuo sin reuniones formales.'],
    ],
    hdr_color='7B3F9E'
)

heading2('5.4 Hacer Menos (Less Of)')
body('Prácticas que existen y generan fricción; reducir su frecuencia o alcance.', italic=True)
simple_table(
    ['#', 'Práctica a Reducir', 'Por qué y cómo reducirla'],
    [
        ['1', 'Agrupar demasiados módulos en una sola historia de usuario', 'HU-09 tenía 5 módulos y generó una semana 9 muy estresante. En Sprint 3 ninguna HU debe tener más de 3 módulos distintos.'],
        ['2', 'Estimaciones basadas en "parece simple"', 'La integración SUNAT pareció simple y tomó el doble de lo estimado. Agregar siempre un 30% de margen en tareas de integración con servicios externos.'],
    ],
    hdr_color='E67E22'
)

heading2('5.5 Dejar de Hacer (Stop Doing)')
body('Prácticas que deben eliminarse por completo porque generan daño o desperdicio.', italic=True)
simple_table(
    ['#', 'Práctica a Eliminar', 'Alternativa propuesta'],
    [
        ['1', 'Desarrollar sin ningún test automatizado', 'Incorporar al menos un test de integración por historia antes de cerrarla. La calidad no puede depender solo de pruebas manuales con Postman.'],
        ['2', 'Dejar el feedback del cliente solo para el Sprint Review', 'El propietario debe tener acceso continuo al sistema en pruebas. Si algo no le gusta, es mejor saberlo en la semana 2 y no en la semana 5.'],
    ],
    hdr_color='C0392B'
)

doc.add_page_break()

# ══════════════════════════════════════════════════════════════════════════════
# 6. PLAN DE ACCIÓN PARA EL SPRINT 3
# ══════════════════════════════════════════════════════════════════════════════
heading1('6. Plan de Acción para el Sprint 3')
body('La retrospectiva no tiene valor si no genera acciones concretas. Se seleccionaron las mejoras más impactantes y se formularon como compromisos medibles con criterio de éxito claro. Se recomienda entre 3 y 5 acciones prioritarias para asegurar que realmente se implementen.')

heading2('6.1 Acciones de Mejora Comprometidas')
simple_table(
    ['#', 'Acción de Mejora', 'Responsable', 'Prioridad', 'Fecha Límite', 'Métrica de Éxito', 'Estado'],
    [
        ['1', 'Escribir tests de integración para 3 flujos críticos (venta, boleta, cierre de caja)',
         'Alumno dev', 'ALTA', 'Semana 10', 'Cobertura ≥ 50% en VentaService, FacturacionService y CajaService', 'Pendiente'],
        ['2', 'Dividir HUs del Sprint 3 con máximo 3 módulos por historia',
         'Alumno dev', 'ALTA', 'Sprint Planning S3', 'Ninguna semana supera las 22h de trabajo', 'Pendiente'],
        ['3', 'Configurar acceso del propietario al sistema de pruebas desde inicio del sprint',
         'Alumno dev', 'ALTA', 'Semana 10', 'El propietario envía al menos 1 feedback intermedio antes de la semana 13', 'Pendiente'],
        ['4', 'Implementar el XML del RC real para producción (Resumen Diario de Boletas)',
         'Alumno dev', 'MEDIA', 'Semana 12', 'RC enviado y validado en sandbox con al menos 3 boletas agrupadas', 'Pendiente'],
        ['5', 'Probar cada pantalla Android en dispositivo físico el mismo día que se implementa',
         'Alumno dev', 'MEDIA', 'Sprint 3 (continuo)', '0 errores de UI detectados en la demo final de Android', 'Pendiente'],
        ['6', 'Agregar 30% de margen a tareas con servicios externos (IA, Android)',
         'Alumno dev', 'BAJA', 'Sprint Planning S3', 'Horas reales ≤ horas estimadas al cierre del Sprint 3', 'Pendiente'],
    ]
)

heading2('6.2 Tareas Pendientes con Deuda Técnica: Estrategia para Sprint 3')
body('A diferencia del Sprint 1, el Sprint 2 no dejó historias incompletas. Sin embargo, se identificaron elementos de deuda técnica que deben atenderse:')
simple_table(
    ['Elemento', 'Tipo', 'Recomendación para Sprint 3'],
    [
        ['RCB real (XML RC agrupado de boletas)', 'Deuda técnica — funcionalidad incompleta', 'Prioridad media. Implementar antes de la puesta en producción. Estimación: 3 SP adicionales en Sprint 3.'],
        ['Pruebas automatizadas (0% cobertura actual)', 'Deuda técnica — calidad', 'Tarea transversal en Sprint 3. Objetivo: tests para VentaService, InventarioService y FacturacionService.'],
        ['Desglose por vendedor en cierre de caja', 'Ajuste solicitado por el propietario', 'Prioridad baja. Incluir como tarea menor en la HU de reportes del Sprint 3.'],
    ]
)

doc.add_page_break()

# ══════════════════════════════════════════════════════════════════════════════
# 7. ANÁLISIS DE VELOCITY Y PROYECCIÓN
# ══════════════════════════════════════════════════════════════════════════════
heading1('7. Análisis de Velocity y Proyección')
body('La velocity es la cantidad de story points que el desarrollador completa en un sprint. Con dos sprints completos disponibles, la proyección del Sprint 3 es significativamente más confiable que al inicio del proyecto.')

heading2('7.1 Velocity del Sprint 2')
simple_table(
    ['Concepto', 'Valor'],
    [
        ['Velocity planificada',               '32 Story Points'],
        ['Velocity real (SP completados)',      '32 Story Points'],
        ['Ratio de cumplimiento',               '100%'],
        ['Velocity acumulada (Sprint 1 + 2)',   '52 SP en 9 semanas'],
        ['Velocity promedio por semana',        '5.8 SP / semana'],
        ['Velocity recomendada para Sprint 3',  '34–38 Story Points'],
    ]
)

heading2('7.2 Proyección del Proyecto')
body('Con dos sprints de datos reales, el proyecto muestra una tendencia positiva: la velocity aumentó de 20 SP (Sprint 1, 4 semanas) a 32 SP (Sprint 2, 5 semanas), lo que refleja que el desarrollador gana confianza con el stack tecnológico y la arquitectura del sistema.')
simple_table(
    ['Sprint', 'Velocity Real/Estimada', 'SP Acumulados', 'Entregable Principal'],
    [
        ['Sprint 1 (real)',      '20 SP', '20 SP', 'Base del sistema: productos, usuarios, auth, BD'],
        ['Sprint 2 (real)',      '32 SP', '52 SP', 'Sistema operativo: POS, SUNAT, inventario, caja, gastos, devoluciones, crédito'],
        ['Sprint 3 (estimado)', '36 SP', '88 SP', 'Funciones avanzadas: promociones, fidelización, reportes con IA, app Android'],
        ['Despliegue (est.)',    '—',     '—',     'Puesta en producción, capacitación, documentación final'],
    ]
)

doc.add_page_break()

# ══════════════════════════════════════════════════════════════════════════════
# 8. DEFINICIÓN DE TERMINADO (DoD)
# ══════════════════════════════════════════════════════════════════════════════
heading1('8. Revisión de la Definición de Terminado (DoD)')
body('La Definición de Terminado es el acuerdo sobre qué criterios debe cumplir una historia para considerarse completada. Tras la experiencia del Sprint 2, se propone actualizar el DoD para el Sprint 3 incorporando las lecciones aprendidas.')

heading2('8.1 DoD Sprint 2 vs DoD Propuesto Sprint 3')
simple_table(
    ['DoD Sprint 2 (aplicado)', 'DoD Sprint 3 (mejorado)'],
    [
        ['1. Código desarrollado y funcional',
         '1. Código desarrollado y funcional'],
        ['2. Endpoints probados con Postman (pruebas manuales)',
         '2. Endpoints probados con Postman y al menos 1 test automatizado de integración'],
        ['3. Frontend conectado al backend y funcional',
         '3. Frontend conectado al backend y funcional'],
        ['4. Stock / datos actualizados correctamente tras cada operación',
         '4. Stock / datos actualizados correctamente tras cada operación'],
        ['5. Migración Flyway ejecutada sin errores (si aplica)',
         '5. Migración Flyway ejecutada sin errores (si aplica)'],
        ['6. Funcionalidad demostrada al cliente',
         '6. Criterios de aceptación verificados por checklist previamente definida'],
        ['(sin criterio para Android)',
         '7. Pantalla Android probada en dispositivo físico (si el módulo incluye Android)'],
        ['(sin criterio de performance)',
         '8. Respuesta del endpoint < 500 ms en condiciones normales (medido en Postman)'],
        ['(sin criterio de aprobación formal)',
         '9. Demo aprobada por el propietario antes de cerrar la historia'],
    ],
    hdr_color='1E408A'
)
body('Los cambios principales son la incorporación de tests automatizados, un checklist de criterios de aceptación por historia, pruebas en dispositivo físico para Android, validación de performance y aprobación formal del propietario. Estos cambios buscan que el Sprint 3 llegue al despliegue con mayor confianza técnica y menos riesgo de regresión.')

doc.add_page_break()

# ══════════════════════════════════════════════════════════════════════════════
# 9. INDICADORES DE SALUD
# ══════════════════════════════════════════════════════════════════════════════
heading1('9. Indicadores de Salud del Equipo')
body('Más allá de las métricas del producto, es importante evaluar cómo se sintió el equipo durante el sprint. En este proyecto el "equipo" incluye al desarrollador y al propietario como usuario principal. Se evaluaron los siguientes indicadores de 1 a 5.')
simple_table(
    ['Indicador', 'Puntuación', 'Comentario del Equipo'],
    [
        ['Claridad del objetivo del sprint',  '4.5 / 5', 'El objetivo era muy concreto: que el negocio pueda operar digitalmente. No hubo ambigüedad en qué construir.'],
        ['Colaboración con el cliente (propietario)', '4.0 / 5', 'Buena disponibilidad para las demos. Se podría mejorar dándole acceso continuo al sistema durante el sprint.'],
        ['Calidad del Sprint Planning', '3.5 / 5', 'Mejoró respecto al Sprint 1, pero HU-09 quedó demasiado cargada. La estimación de SUNAT fue optimista.'],
        ['Nivel de estrés', '3.2 / 5', 'La semana 9 fue la más intensa del proyecto. Cinco módulos en una semana para un solo desarrollador es demasiado.'],
        ['Confianza en la entrega', '4.5 / 5', 'Se logró completar el 100% del sprint. El sistema funciona de punta a punta y el propietario quedó satisfecho.'],
        ['Satisfacción general del sprint', '4.3 / 5', 'Es el sprint más importante del proyecto: el negocio ya puede operar con el sistema. Eso genera mucha motivación.'],
    ]
)
body('La puntuación promedio general es de 4.0 / 5, la más alta de los dos sprints realizados. Los puntos más bajos (calidad del planning con 3.5 y nivel de estrés con 3.2) están directamente relacionados con la densidad de HU-09 y la complejidad subestimada de SUNAT, ambos abordados en el plan de acción.')

doc.add_page_break()

# ══════════════════════════════════════════════════════════════════════════════
# 10. CONCLUSIONES Y COMPROMISOS
# ══════════════════════════════════════════════════════════════════════════════
heading1('10. Conclusiones y Compromisos')

heading2('10.1 Resumen Ejecutivo')
body('El Sprint 2 del Sistema de Gestión Licorería Chilalo Shot fue el sprint más importante del proyecto y se completó al 100%. Las 4 historias de usuario, las 23 tareas técnicas y los 32 story points planificados fueron entregados en las 5 semanas previstas. Al cierre del sprint, el negocio ya puede operar digitalmente: registrar ventas en el POS en menos de 45 segundos, emitir boletas y facturas electrónicas en el ambiente de pruebas de SUNAT, controlar el inventario en tiempo real y gestionar caja, gastos, devoluciones y crédito.')
body('Los principales aprendizajes son: que las integraciones con servicios externos deben estimarse con margen, que las historias de usuario muy densas generan semanas insostenibles, y que la ausencia de pruebas automatizadas representa una deuda técnica que crece con cada sprint. Todos estos puntos tienen acciones concretas comprometidas para el Sprint 3.')

heading2('10.2 Los 3 Compromisos Principales del Equipo')
numbered('Compromiso 1: Escribir tests de integración automatizados para los tres flujos críticos del sistema (crear venta, emitir boleta y calcular cierre de caja) antes de finalizar la semana 10, como primera actividad del Sprint 3.')
numbered('Compromiso 2: Dividir el Sprint 3 de manera que ninguna historia de usuario agrupe más de 3 módulos distintos, evitando repetir la semana 9 del Sprint 2 donde se concentraron 5 módulos simultáneos.')
numbered('Compromiso 3: Configurar el acceso del propietario al sistema de pruebas desde el inicio del Sprint 3, para recibir feedback continuo sobre promociones, reportes y la app Android, y no solo al final en el Sprint Review.')

heading2('10.3 Seguimiento')
body('Estos compromisos serán revisados en la retrospectiva del Sprint 3 para verificar su implementación. El alumno desarrollador los registrará como las primeras tareas del Sprint 3 y los comunicará al propietario y al docente evaluador al inicio de la semana 10.')

doc.add_paragraph()
doc.add_paragraph()

# Firmas
t_firma = doc.add_table(rows=2, cols=2)
t_firma.style = 'Table Grid'
for r in t_firma.rows:
    for c in r.cells:
        set_cell_bg(c, 'FFFFFF')

# líneas de firma
for ci in range(2):
    cell = t_firma.rows[0].cells[ci]
    set_cell_border(cell, bottom='1E408A')
    cell.paragraphs[0].add_run('   ')

labels = ['Alumno desarrollador (Scrum Master / Dev)', 'Propietario Chilalo Shot (Product Owner)']
for ci in range(2):
    cell = t_firma.rows[1].cells[ci]
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(labels[ci])
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor(0x70,0x70,0x70)

doc.add_paragraph()
p_fecha = doc.add_paragraph()
p_fecha.alignment = WD_ALIGN_PARAGRAPH.CENTER
rf = p_fecha.add_run('Fecha: Fin de semana 9 — Marzo 2025')
rf.italic = True
rf.font.size = Pt(9)
rf.font.color.rgb = RGBColor(0x70,0x70,0x70)

doc.add_paragraph()
p_pie = doc.add_paragraph()
p_pie.alignment = WD_ALIGN_PARAGRAPH.CENTER
rp = p_pie.add_run('Documento elaborado como parte del Trabajo Académico Aplicado (TAA). Instituto de Educación Superior – Piura, Perú.')
rp.italic = True
rp.font.size = Pt(8)
rp.font.color.rgb = RGBColor(0x90,0x90,0x90)

# ══════════════════════════════════════════════════════════════════════════════
# GUARDAR
# ══════════════════════════════════════════════════════════════════════════════
out = r'C:\Users\Usuario\Desktop\Liquor-Store-Piura-master\Liquor-Store-Piura-master\Retrospectiva_Sprint2.docx'
doc.save(out)
print(f'Guardado: {out}')
