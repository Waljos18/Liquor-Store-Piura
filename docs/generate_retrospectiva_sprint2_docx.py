from docx import Document
from docx.shared import Pt


def add_title(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(16)


def add_h1(doc: Document, text: str) -> None:
    doc.add_heading(text, level=1)


def add_h2(doc: Document, text: str) -> None:
    doc.add_heading(text, level=2)


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def add_table(doc: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = h
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = value


def main() -> None:
    doc = Document()

    add_title(doc, "RETROSPECTIVA DEL SPRINT")
    doc.add_paragraph("Sprint 2 - Sistema de Gestion para Licoreria Chilalo Shot")
    doc.add_paragraph("Proyecto: Sistema Web y App Movil - Chilalo Shot")
    doc.add_paragraph("Sprint: Sprint 2")
    doc.add_paragraph("Fecha: 07 de abril de 2026")
    doc.add_paragraph("Facilitador: Equipo de desarrollo")
    doc.add_paragraph("Participantes: Desarrollador, propietario del negocio y docente evaluador")
    doc.add_paragraph("Duracion: 1 hora")
    doc.add_paragraph("Documento de uso interno y academico")

    add_h1(doc, "Tabla de Contenidos")
    toc_items = [
        "Tabla de Contenidos",
        "1. Introduccion: Que es una Retrospectiva del Sprint?",
        "1.1 Principios Fundamentales",
        "1.2 Formato de esta Retrospectiva",
        "2. Contexto del Sprint 2",
        "2.1 Metricas Clave del Sprint 2",
        "2.2 Estado de las Historias de Usuario",
        "2.3 Interpretacion de los Datos",
        "3. Que Salio Bien? (Glad)",
        "4. Que No Salio Bien? (Sad / Mad)",
        "4.1 Problemas Identificados",
        "4.2 Diagrama de Causa y Efecto",
        "5. Analisis Starfish (Estrella de Mar)",
        "5.1 Seguir Haciendo (Keep Doing)",
        "5.2 Hacer Mas (More Of)",
        "5.3 Empezar a Hacer (Start Doing)",
        "5.4 Hacer Menos (Less Of)",
        "5.5 Dejar de Hacer (Stop Doing)",
        "6. Plan de Accion para el Sprint 3",
        "6.1 Acciones de Mejora Comprometidas",
        "6.2 Historias No Completadas: Estrategia para Sprint 3",
        "7. Analisis de Velocity y Proyeccion",
        "7.1 Velocity del Sprint 2",
        "7.2 Proyeccion del Proyecto",
        "8. Revision de la Definicion de Terminado (DoD)",
        "8.1 DoD Sprint 2 vs DoD Propuesto Sprint 3",
        "9. Indicadores de Salud del Equipo",
        "10. Conclusiones y Compromisos",
        "10.1 Resumen Ejecutivo",
        "10.2 Los 3 Compromisos Principales del Equipo",
        "10.3 Seguimiento",
    ]
    for item in toc_items:
        doc.add_paragraph(item)

    add_h1(doc, "1. Introduccion: Que es una Retrospectiva del Sprint?")
    doc.add_paragraph(
        "La retrospectiva permite revisar como trabajamos durante el sprint, identificar aciertos, "
        "aprender de las dificultades y acordar mejoras concretas para el siguiente ciclo."
    )
    add_h2(doc, "1.1 Principios Fundamentales")
    add_bullets(
        doc,
        [
            "Transparencia: hablar con claridad sobre resultados y dificultades.",
            "Inspeccion: usar datos reales para evaluar el sprint.",
            "Adaptacion: definir acciones concretas para mejorar en Sprint 3.",
        ],
    )
    add_h2(doc, "1.2 Formato de esta Retrospectiva")
    add_bullets(
        doc,
        [
            "Analisis cuantitativo de metricas del Sprint 2.",
            "Revision de logros y problemas (Glad / Sad / Mad).",
            "Analisis Starfish para acciones de mejora.",
            "Plan de accion con responsables y resultados esperados.",
        ],
    )

    add_h1(doc, "2. Contexto del Sprint 2")
    doc.add_paragraph(
        "El Sprint 2 se enfoco en el nucleo operativo de la licoreria: ventas en POS, "
        "inventario, compras, caja, gastos, devoluciones, credito y comprobantes electronicos."
    )
    add_h2(doc, "2.1 Metricas Clave del Sprint 2")
    add_table(
        doc,
        ["Metrica", "Planificado", "Real", "Cumplimiento"],
        [
            ["Story Points comprometidos", "139 SP", "139 SP", "100%"],
            ["Story Points completados (Velocity)", "139 SP", "139 SP", "100%"],
            ["Historias de Usuario completadas", "4", "4 de 4", "100%"],
            ["Tareas completadas", "23", "23 de 23", "100%"],
            ["Horas aproximadas", "~100 h", "~100 h", "En rango"],
            ["Duracion del Sprint", "5 semanas", "5 semanas", "100%"],
        ],
    )
    add_h2(doc, "2.2 Estado de las Historias de Usuario")
    add_table(
        doc,
        ["ID", "Historia (resumen)", "Estado"],
        [
            ["HU-06", "POS web para ventas rapidas con multiples pagos", "Completada"],
            ["HU-07", "Comprobantes electronicos para la operacion", "Completada"],
            ["HU-08", "Control de inventario, compras y mermas", "Completada"],
            ["HU-09", "Caja, gastos, devoluciones, credito y clientes", "Completada"],
        ],
    )
    add_h2(doc, "2.3 Interpretacion de los Datos")
    doc.add_paragraph(
        "El Sprint 2 tuvo cumplimiento total del backlog. Se logro implementar el flujo operativo "
        "diario del negocio de extremo a extremo, incluyendo control de stock, ventas y soporte "
        "administrativo para una gestion mas ordenada."
    )

    add_h1(doc, "3. Que Salio Bien? (Glad)")
    add_table(
        doc,
        ["#", "Aspecto positivo", "Evidencia"],
        [
            ["1", "Cobertura operativa completa", "Se cubrieron ventas, inventario, compras y caja."],
            ["2", "Cumplimiento del backlog", "23 de 23 tareas finalizadas."],
            ["3", "Mejora visible para el cliente", "Operacion diaria digital y trazable."],
            ["4", "Respuesta a feedback previo", "Stock visible en flujo de ventas."],
            ["5", "Buena coordinacion en cierre", "Sin historias pendientes al final del sprint."],
        ],
    )

    add_h1(doc, "4. Que No Salio Bien? (Sad / Mad)")
    add_h2(doc, "4.1 Problemas Identificados")
    add_table(
        doc,
        ["#", "Nivel", "Problema", "Causa raiz", "Impacto"],
        [
            [
                "1",
                "SAD",
                "Ajustes en comprobantes electronicos en pruebas",
                "Variaciones de validacion en integracion externa",
                "Retrabajo tecnico puntual en fase final",
            ],
            [
                "2",
                "SAD",
                "Sobrecarga de validaciones al cierre",
                "Muchos modulos criticos convergieron al final",
                "Mayor presion en la semana final",
            ],
            [
                "3",
                "MAD",
                "Reportes de avance con texto tecnico",
                "Documentacion pensada mas para equipo tecnico",
                "Menor claridad para presentacion a cliente",
            ],
        ],
    )
    add_h2(doc, "4.2 Diagrama de Causa y Efecto")
    add_bullets(
        doc,
        [
            "Causa: validaciones externas de comprobantes -> Efecto: ajustes adicionales en cierre.",
            "Causa: alta concentracion de modulos criticos -> Efecto: mayor carga final.",
            "Causa: enfoque tecnico en redaccion -> Efecto: explicacion menos simple al cliente.",
        ],
    )

    add_h1(doc, "5. Analisis Starfish (Estrella de Mar)")
    add_h2(doc, "5.1 Seguir Haciendo (Keep Doing)")
    add_bullets(
        doc,
        [
            "Priorizar funcionalidades con impacto diario en el negocio.",
            "Validar entregas con escenarios reales de uso.",
            "Mantener seguimiento de backlog por historia.",
        ],
    )
    add_h2(doc, "5.2 Hacer Mas (More Of)")
    add_bullets(
        doc,
        [
            "Demostraciones intermedias por modulo durante el sprint.",
            "Pruebas anticipadas de integraciones externas.",
            "Revision de textos orientados a lenguaje cliente.",
        ],
    )
    add_h2(doc, "5.3 Empezar a Hacer (Start Doing)")
    add_bullets(
        doc,
        [
            "Checklist de cierre por modulo una semana antes del fin del sprint.",
            "Guion de demo ejecutiva para cliente no tecnico.",
            "Registro de riesgos por integraciones externas desde el inicio.",
        ],
    )
    add_h2(doc, "5.4 Hacer Menos (Less Of)")
    add_bullets(
        doc,
        [
            "Concentrar validaciones complejas en la ultima semana.",
            "Explicaciones largas con terminologia tecnica en review.",
        ],
    )
    add_h2(doc, "5.5 Dejar de Hacer (Stop Doing)")
    add_bullets(
        doc,
        [
            "Posponer pruebas de integracion para el cierre.",
            "Presentar resultados sin convertirlos a valor de negocio.",
        ],
    )

    add_h1(doc, "6. Plan de Accion para el Sprint 3")
    add_h2(doc, "6.1 Acciones de Mejora Comprometidas")
    add_table(
        doc,
        ["#", "Accion", "Responsable", "Prioridad", "Fecha limite", "Metrica de exito"],
        [
            [
                "1",
                "Realizar pruebas de integraciones desde la semana 1",
                "Desarrollador",
                "Alta",
                "Inicio Sprint 3",
                "0 bloqueos criticos de integracion en cierre",
            ],
            [
                "2",
                "Estandarizar lenguaje cliente en reportes y demos",
                "Equipo",
                "Alta",
                "Todo Sprint 3",
                "100% de historias explicadas por valor de negocio",
            ],
            [
                "3",
                "Aplicar checklist de cierre por modulo",
                "Equipo",
                "Media",
                "Ultima semana Sprint 3",
                "Menos retrabajo en validacion final",
            ],
        ],
    )
    add_h2(doc, "6.2 Historias No Completadas: Estrategia para Sprint 3")
    add_table(
        doc,
        ["Historia", "Estado", "Estrategia para Sprint 3"],
        [
            ["Ninguna historia pendiente", "Completado", "No aplica; Sprint 2 cerro al 100%."],
            [
                "Nuevas prioridades del cliente",
                "Backlog futuro",
                "Enfocar Sprint 3 en promociones, reportes e inteligencia de negocio.",
            ],
        ],
    )

    add_h1(doc, "7. Analisis de Velocity y Proyeccion")
    add_h2(doc, "7.1 Velocity del Sprint 2")
    add_table(
        doc,
        ["Concepto", "Valor"],
        [
            ["Velocity planificada", "139 SP"],
            ["Velocity real", "139 SP"],
            ["Ratio de cumplimiento", "100%"],
            ["Lectura", "Capacidad alineada con planificacion"],
        ],
    )
    add_h2(doc, "7.2 Proyeccion del Proyecto")
    doc.add_paragraph(
        "Con Sprint 2 completado al 100%, Sprint 3 puede enfocarse en funcionalidades de valor "
        "diferencial: promociones, fidelizacion, app movil, reportes y cierre de implementacion."
    )

    add_h1(doc, "8. Revision de la Definicion de Terminado (DoD)")
    add_h2(doc, "8.1 DoD Sprint 2 vs DoD Propuesto Sprint 3")
    add_table(
        doc,
        ["DoD Sprint 2", "DoD Propuesto Sprint 3"],
        [
            [
                "Funcionalidad completa y probada en flujo principal",
                "Funcionalidad completa + evidencia de valor para cliente",
            ],
            [
                "Validacion tecnica del modulo",
                "Validacion tecnica + pruebas en escenarios reales de negocio",
            ],
            [
                "Demostracion en review",
                "Demostracion con guion ejecutivo y lenguaje no tecnico",
            ],
        ],
    )

    add_h1(doc, "9. Indicadores de Salud del Equipo")
    add_table(
        doc,
        ["Indicador", "Puntuacion (1-5)", "Comentario"],
        [
            ["Claridad del objetivo", "4.9", "El enfoque operativo fue claro desde el inicio."],
            ["Colaboracion", "4.8", "Buena coordinacion para cerrar modulos interdependientes."],
            ["Ritmo de trabajo", "4.4", "Ritmo alto, con presion en el cierre."],
            ["Confianza en la entrega", "4.9", "Backlog completado en su totalidad."],
            ["Satisfaccion general", "4.8", "Resultado fuerte y alineado al negocio."],
        ],
    )

    add_h1(doc, "10. Conclusiones y Compromisos")
    add_h2(doc, "10.1 Resumen Ejecutivo")
    doc.add_paragraph(
        "El Sprint 2 fue exitoso y de alto impacto: se digitalizo el flujo operativo diario del negocio "
        "y se cumplio el 100% del backlog. El sistema quedo preparado para pasar de operacion base "
        "a capacidades diferenciales en Sprint 3."
    )
    add_h2(doc, "10.2 Los 3 Compromisos Principales del Equipo")
    add_bullets(
        doc,
        [
            "Anticipar validaciones de integraciones externas desde el inicio.",
            "Mantener comunicacion orientada al valor del negocio en todas las demos.",
            "Reducir carga de cierre con checklist progresivo por modulo.",
        ],
    )
    add_h2(doc, "10.3 Seguimiento")
    doc.add_paragraph(
        "Los compromisos se revisaran en la retrospectiva del Sprint 3 con evidencia de cumplimiento "
        "en demos, reportes y validacion del cliente."
    )

    out = (
        r"c:\Users\Usuario\Desktop\Liquor-Store-Piura-master\Liquor-Store-Piura-master"
        r"\Retrospectiva_Sprint2_CHILALO_SHOT.docx"
    )
    doc.save(out)
    print(out)


if __name__ == "__main__":
    main()
